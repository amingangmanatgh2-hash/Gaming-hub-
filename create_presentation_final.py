from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageEnhance
from lxml import etree
import xml.etree.ElementTree as ET
import os, math, wave, struct, zipfile

OUT='output/چرا_به_قانون_نیاز_داریم_نسخه_نهایی.pptx'
A='assets_v3'; V='assets_final'; os.makedirs(V,exist_ok=True); os.makedirs('output',exist_ok=True)
SW,SH=13.333,7.5
INK=(3,10,17); NAVY=(6,22,36); PANEL=(8,29,46); WHITE=(247,250,252); ICE=(201,229,241)
GOLD=(241,183,62); BLUE=(69,162,224); CYAN=(80,205,215); MUTED=(126,153,169); CORAL=(242,108,93)
FONT='Tahoma'; FA=str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')

src={'justice':f'{A}/hero_justice.png','school':f'{A}/hero_school_clean.png','city':f'{A}/hero_city_clean.png','traffic':f'{A}/traffic_3d.png','camp':f'{A}/camp_3d.png','constitution':f'{A}/constitution_3d.png'}
def prep(path,out,dark='right'):
    im=Image.open(path).convert('RGB'); W,H=1600,900; sc=max(W/im.width,H/im.height)
    im=im.resize((int(im.width*sc),int(im.height*sc)),Image.Resampling.LANCZOS)
    im=im.crop(((im.width-W)//2,(im.height-H)//2,(im.width+W)//2,(im.height+H)//2)); im=ImageEnhance.Contrast(im).enhance(1.06)
    ov=Image.new('RGBA',(W,H)); d=ImageDraw.Draw(ov)
    for x in range(W):
        t=x/(W-1); a=int(228*(t**1.35 if dark=='right' else (1-t)**1.35)); d.line((x,0,x,H),fill=(2,8,14,a))
    Image.alpha_composite(im.convert('RGBA'),ov).convert('RGB').save(out,'JPEG',quality=91,optimize=True,progressive=True)
for k,p in src.items():
    prep(p,f'{V}/{k}_right.jpg','right'); prep(p,f'{V}/{k}_left.jpg','left')

# Tiny transparent animated pulse: visual motion without a heavy full-frame GIF.
def pulse_gif(path):
    frames=[]; N=28; S=180
    for i in range(N):
        im=Image.new('RGBA',(S,S),(0,0,0,0)); d=ImageDraw.Draw(im); q=(math.sin(i/N*2*math.pi)+1)/2
        r=38+26*q; alpha=int(210-105*q); w=5
        d.ellipse((S/2-r,S/2-r,S/2+r,S/2+r),outline=(*GOLD,alpha),width=w)
        d.ellipse((S/2-10,S/2-10,S/2+10,S/2+10),fill=(*BLUE,220))
        frames.append(im)
    frames[0].save(path,save_all=True,append_images=frames[1:],duration=55,loop=0,disposal=2,optimize=True)
pulse_gif(f'{V}/pulse.gif')

# Two restrained, very short transition sounds.
def make_sound(path,kind):
    sr=22050; dur=.34 if kind=='whoosh' else .25; n=int(sr*dur); raw=[]
    for i in range(n):
        t=i/sr; env=math.sin(math.pi*i/n)**2
        if kind=='whoosh': val=.11*env*math.sin(2*math.pi*(180+520*t)*t)
        else: val=.10*env*(math.sin(2*math.pi*660*t)+.45*math.sin(2*math.pi*990*t))
        raw.append(struct.pack('<h',int(max(-1,min(1,val))*32767)))
    with wave.open(path,'wb') as w: w.setparams((1,2,sr,n,'NONE','not compressed')); w.writeframes(b''.join(raw))
make_sound(f'{V}/whoosh.wav','whoosh'); make_sound(f'{V}/chime.wav','chime')

prs=Presentation(); prs.slide_width=Inches(SW); prs.slide_height=Inches(SH); blank=prs.slide_layouts[6]
def C(v): return RGBColor(*v)
def setname(sh,n):
    try:
        if sh.shape_type==13: sh._element.nvPicPr.cNvPr.set('name',n)
        else: sh._element.nvSpPr.cNvPr.set('name',n)
    except: pass
    return sh
def rtl(p,center=False):
    p.alignment=PP_ALIGN.CENTER if center else PP_ALIGN.RIGHT; pr=p._p.get_or_add_pPr(); pr.set('rtl','1'); pr.set('algn','ctr' if center else 'r')
def text(sl,x,y,w,h,s,fs=22,color=WHITE,bold=False,center=False,top=False,name=None):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.vertical_anchor=MSO_ANCHOR.TOP if top else MSO_ANCHOR.MIDDLE; tf.margin_left=tf.margin_right=Inches(.04); tf.margin_top=tf.margin_bottom=Inches(.02)
    p=tf.paragraphs[0]; p.text=s; rtl(p,center); p.line_spacing=1.05
    for r in p.runs: r.font.name=FONT; r.font.size=Pt(fs); r.font.bold=bold; r.font.color.rgb=C(color); r._r.get_or_add_rPr().set('lang','fa-IR')
    if name: setname(sh,name)
    return sh
def shadow(sh):
    sp=sh._element.spPr; eff=OxmlElement('a:effectLst'); o=OxmlElement('a:outerShdw'); o.set('blurRad','65000'); o.set('dist','25000'); o.set('dir','2700000'); o.set('algn','ctr'); o.set('rotWithShape','0')
    c=OxmlElement('a:srgbClr'); c.set('val','000000'); a=OxmlElement('a:alpha'); a.set('val','45000'); c.append(a); o.append(c); eff.append(o); sp.append(eff)
def shape(sl,kind,x,y,w,h,fill=PANEL,line=None,trans=0,name=None,shad=False):
    sh=sl.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=C(fill); sh.fill.transparency=trans
    if line: sh.line.color.rgb=C(line); sh.line.width=Pt(1.25)
    else: sh.line.fill.background()
    if shad: shadow(sh)
    if name: setname(sh,name)
    return sh
def picture(sl,key,dark='right'):
    p=sl.shapes.add_picture(f'{V}/{key}_{dark}.jpg',0,0,width=prs.slide_width,height=prs.slide_height); setname(p,'!!Scene'); return p
def accent(n):
    if n<=3:return GOLD
    if n<=7:return BLUE
    if n<=11:return CYAN
    if n<=15:return GOLD
    return BLUE
def presenter(n):
    if n in {1,2,3,18}: return 'امین کریم‌پور'
    if 4<=n<=7: return 'عزیزی'
    if 8<=n<=11: return 'طاها عابدی‌فر'
    if 12<=n<=14: return 'گله‌دار'
    return 'محمدرضا فولادی'
def footer(sl,n):
    ac=accent(n); text(sl,.68,7.05,3.0,.22,'مطالعات اجتماعی هفتم • درس ۳',8.5,MUTED)
    text(sl,8.65,7.02,2.55,.22,'ارائه: '+presenter(n),8.5,ac,True)
    text(sl,11.82,6.99,.82,.28,(str(n).zfill(2)+' / 18').translate(FA),10,ac,True,True)
    shape(sl,MSO_SHAPE.RECTANGLE,.68,6.84,11.96,.018,(45,67,80)); shape(sl,MSO_SHAPE.RECTANGLE,.68,6.82,11.96*n/18,.045,ac,name='!!Progress')
def base(n,kicker,title,img=None,dark='right'):
    sl=prs.slides.add_slide(blank); ac=accent(n)
    if img: picture(sl,img,dark)
    else: shape(sl,MSO_SHAPE.RECTANGLE,0,0,SW,SH,INK)
    shape(sl,MSO_SHAPE.RECTANGLE,0,0,SW,1.36,INK,None,23,name='!!TopShade')
    shape(sl,MSO_SHAPE.ROUNDED_RECTANGLE,.68,.28,1.72,.33,NAVY,ac,2,shad=True)
    text(sl,.77,.31,1.54,.23,kicker,9.5,ac,True,True,name='!!Kicker')
    text(sl,.68,.7,11.7,.48,title,29,WHITE,True,name='!!Title')
    shape(sl,MSO_SHAPE.RECTANGLE,.68,1.27,1.45,.026,ac,name='!!Rule')
    footer(sl,n); return sl
def glass(sl,x,y,w,h,col=GOLD):
    shape(sl,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,(6,26,42),(91,122,139),9,shad=True)
    shape(sl,MSO_SHAPE.ROUNDED_RECTANGLE,x+.03,y+.03,w-.06,.05,WHITE,None,86)
    shape(sl,MSO_SHAPE.ROUNDED_RECTANGLE,x+w-.055,y+.2,.035,h-.4,col)
def focus(sl,x,y,w,lead,sub,col=GOLD):
    glass(sl,x-.18,y-.16,w+.36,2.18,col); text(sl,x,y,w,.66,lead,27,col,True); shape(sl,MSO_SHAPE.RECTANGLE,x,y+.72,1.05,.025,col); text(sl,x,y+.9,w,1.0,sub,16.5,ICE,top=True)
def badge(sl,x,y,num,label,col):
    shape(sl,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,2.9,.64,NAVY,col,2,shad=True); text(sl,x+.1,y+.06,.74,.48,num,21,col,True,True); text(sl,x+.95,y+.08,1.76,.4,label,12.5,WHITE,True)
def pulse(sl,x,y):
    sl.shapes.add_picture(f'{V}/pulse.gif',Inches(x),Inches(y),width=Inches(.72),height=Inches(.72))

# 1 — title
s=prs.slides.add_slide(blank); picture(s,'justice','left'); shape(s,MSO_SHAPE.RECTANGLE,.7,.58,.045,5.92,GOLD,name='!!Rail')
text(s,1.0,.69,5.8,.28,'درس ۳ • مطالعات اجتماعی پایه هفتم',10.5,GOLD,True)
text(s,.98,1.42,5.9,1.46,'چرا به قانون\nنیاز داریم؟',38,WHITE,True,name='!!Hero')
text(s,1.0,3.2,5.35,.55,'نظم، امنیت، عدالت',20,ICE)
shape(s,MSO_SHAPE.RECTANGLE,1.0,4.02,4.65,.02,(82,111,128))
text(s,1.0,4.33,5.25,.88,'گروه ۶ • کلاس ۷/۳\nامین کریم‌پور · عزیزی · طاها عابدی‌فر · گله‌دار · محمدرضا فولادی',11.5,WHITE,top=True); footer(s,1)
# 2 — contents
s=base(2,'فهرست','مسیر ده‌دقیقه‌ای ارائه')
chap=[('۰۱','شروع مسئله','اسلاید ۳',GOLD),('۰۲','چرا قانون؟','اسلایدهای ۴ تا ۷',BLUE),('۰۳','زندگی روزمره','اسلایدهای ۸ تا ۱۱',CYAN),('۰۴','حق و قانون اساسی','اسلایدهای ۱۲ تا ۱۵',GOLD),('۰۵','جمع‌بندی','اسلایدهای ۱۶ تا ۱۸',BLUE)]
shape(s,MSO_SHAPE.RECTANGLE,1.26,1.62,.024,4.58,GOLD)
for i,(n,h,r,c) in enumerate(chap):
    y=1.53+i*.93; text(s,.72,y,.48,.48,n,16,c,True,True); text(s,1.55,y,6.8,.42,h,18,WHITE,True); text(s,9.65,y+.04,2.25,.3,r,11,MUTED,True,True); shape(s,MSO_SHAPE.RECTANGLE,1.55,y+.53,9.9-i*.32,.014,c)
# 3 — problem
s=base(3,'شروع مسئله','اگر یک روز هیچ قانونی نبود…','city','left'); pulse(s,11.65,5.55)
focus(s,.88,1.68,5.25,'چه اتفاقی می‌افتاد؟','ترافیک، مدرسه و شهر خیلی زود بی‌نظم می‌شدند؛ چون مرز حق افراد روشن نبود.',GOLD)
badge(s,.9,4.65,'؟','چه کسی از حق ما دفاع می‌کرد؟',GOLD)
# 4-7 — why law
s=base(4,'دلیل اول','قانون از حق افراد محافظت می‌کند','justice','left'); focus(s,.88,1.72,5.22,'زور، جای حق را نمی‌گیرد','همه می‌توانند شکایت کنند و رسیدگی عادلانه بخواهند.',BLUE); badge(s,.9,4.72,'۱','حفظ حقوق',BLUE)
s=base(5,'دلیل دوم','قانون نظم و امنیت می‌سازد','traffic','right'); focus(s,7.12,1.7,5.18,'قاعدهٔ مشترک = پیش‌بینی‌پذیری','وقتی همه یک قانون را رعایت کنند، خطر و آشفتگی کمتر می‌شود.',BLUE); badge(s,9.42,4.7,'۲','نظم و امنیت',BLUE)
s=base(6,'دلیل سوم','قانون مسئولیت‌پذیری را تمرین می‌دهد'); text(s,.9,1.65,11.5,.7,'«قبل از عمل، پیامد آن را ببین.»',28,BLUE,True,True)
for i,(h,b) in enumerate([('انتخاب','رفتار را آگاهانه انتخاب می‌کنیم'),('پیامد','نتیجهٔ تصمیم را می‌سنجیم'),('پذیرش','مسئولیت کار خود را می‌پذیریم')]):
    x=.98+i*4.1; glass(s,x,3.0,3.55,2.05,BLUE); text(s,x+.2,3.32,3.12,.42,h,18,BLUE,True,True); text(s,x+.24,4.02,3.02,.65,b,13.5,ICE,False,True)
s=base(7,'دلیل چهارم','قانون اختلاف را عادلانه حل می‌کند','justice','right'); focus(s,7.1,1.7,5.2,'دلیل و دادگاه، نه دعوا','قانون به جای زور و انتقام، راه گفت‌وگو و داوری عادلانه را می‌گذارد.',BLUE); badge(s,9.42,4.7,'۴','حل اختلاف',BLUE)
# 8-11 — daily life
s=base(8,'زندگی روزمره','خیابان؛ چند ثانیه عجله، چند حق از دست‌رفته','traffic','right'); pulse(s,.55,5.48); focus(s,7.08,1.7,5.2,'چراغ قرمز قرارداد حفظ جان است','عبور غیرمجاز، جان و زمان دیگران را هم تهدید می‌کند.',CYAN); badge(s,7.12,4.66,'جان','امنیت',CYAN)
s=base(9,'زندگی روزمره','مدرسه؛ نظم برای یادگیری است','school','right'); focus(s,7.08,1.7,5.2,'قانون، کلاس را آرام می‌کند','نوبت صحبت، حضور به‌موقع و احترام؛ فرصت یادگیری را برای همه حفظ می‌کند.',CYAN); badge(s,9.42,4.7,'همه','فرصت برابر',CYAN)
s=base(10,'زندگی روزمره','شهر؛ مال عمومی، مال همه است','city','left'); focus(s,.88,1.7,5.22,'پارک • اتوبوس • کتابخانه','مراقبت از امکانات عمومی یعنی احترام به حق امروز و فردای همه.',CYAN); badge(s,.9,4.7,'ما','امانت‌داری',CYAN)
s=base(11,'زندگی روزمره','اردو؛ قانون خوب باید قابل اجرا باشد','camp','right'); focus(s,7.08,1.7,5.2,'روشن، عادلانه، شدنی','محل وسایل، ساعت استراحت، ایمنی آتش، صف غذا و تقسیم کار باید مشخص باشد.',CYAN); badge(s,9.42,4.72,'۵','قاعدهٔ روشن',CYAN)
# 12-15 — rights and constitution
s=base(12,'حق','حق یعنی چه چیزی باید داشته باشیم؟','school','left'); pulse(s,11.65,5.5)
for i,(h,b,c) in enumerate([('آموزش','فرصت یادگیری',GOLD),('امنیت','زندگی بدون ترس',BLUE),('احترام','حفظ شأن انسان',CYAN)]):
    x=.9+i*4.08; glass(s,x,2.0,3.55,2.55,c); text(s,x+.2,2.4,3.15,.48,h,22,c,True,True); text(s,x+.22,3.35,3.1,.48,b,14.5,WHITE,True,True)
text(s,2.0,5.35,9.3,.5,'قانون، این حق‌ها را برای همه قابل دفاع می‌کند.',17,ICE,True,True)
s=base(13,'تکلیف','هر حق، یک مسئولیت همراه دارد'); text(s,.85,1.7,11.6,.65,'حق من  ↔  مسئولیت من',29,GOLD,True,True)
for i,(a,b) in enumerate([('حق استفاده از کلاس','تکلیف حفظ نظم'),('حق استفاده از پارک','تکلیف مراقبت از وسایل'),('حق احترام','تکلیف احترام به دیگران')]):
    y=2.85+i*1.05; text(s,7.05,y,4.6,.4,a,15.5,GOLD,True); text(s,1.55,y,4.65,.4,b,15.5,BLUE,True); shape(s,MSO_SHAPE.RECTANGLE,6.48,y-.03,.025,.48,MUTED)
s=base(14,'تعادل','حق و تکلیف؛ دو کفهٔ یک ترازو','justice','right'); focus(s,7.1,1.62,5.18,'تعادل جامعه','اگر فقط حق بخواهیم و تکلیف را انجام ندهیم، حق دیگران ضایع می‌شود.',GOLD); badge(s,9.42,4.65,'↔','حق + تکلیف',GOLD)
s=base(15,'قانون اساسی','مادر قوانین کشور','constitution','left'); focus(s,.88,1.62,5.22,'چارچوب اصلی کشور','حقوق عمومی و ساختار نهادها را مشخص می‌کند؛ قوانین دیگر نباید با آن ناسازگار باشند.',GOLD); badge(s,.9,4.68,'۱۲','آذر ۱۳۵۸',GOLD)
# 16-18 — wrap
s=base(16,'مرور','سه سؤال پیش از هر تصمیم'); qs=['آیا کسی آسیب می‌بیند؟','اگر همه انجام دهند چه می‌شود؟','آیا پیامدش را می‌پذیرم؟']
for i,q in enumerate(qs):
    y=1.72+i*1.38; shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,1.38,y,10.55,.82,NAVY,BLUE,2,shad=True); text(s,1.65,y+.1,.72,.58,str(i+1).translate(FA),21,BLUE,True,True); text(s,2.65,y+.11,8.75,.52,q,19,WHITE,True)
s=base(17,'جمع‌بندی','سه نکته‌ای که باید بماند'); pts=[('۱','قانون از حق‌ها محافظت می‌کند.',GOLD),('۲','آزادی با مسئولیت کامل می‌شود.',BLUE),('۳','حق و تکلیف از هم جدا نیستند.',CYAN)]
for i,(n,t,c) in enumerate(pts):
    y=1.72+i*1.38; text(s,1.45,y,.65,.52,n,22,c,True,True); text(s,2.45,y,9.2,.52,t,21,WHITE,True); shape(s,MSO_SHAPE.RECTANGLE,2.45,y+.7,7.9,.018,c)
s=prs.slides.add_slide(blank); picture(s,'justice','left'); shape(s,MSO_SHAPE.RECTANGLE,.7,.58,.045,5.92,GOLD,name='!!Rail'); pulse(s,5.65,5.55)
text(s,1.0,.72,4.8,.25,'پیام پایانی',10.5,GOLD,True); text(s,.98,1.5,5.9,1.52,'قانون یعنی\nاحترام به حق خود و دیگران.',31,WHITE,True,name='!!Hero'); text(s,1.0,3.65,5.25,.7,'جامعهٔ بهتر، از انتخاب مسئولانهٔ ما آغاز می‌شود.',16.5,ICE); text(s,1.0,5.05,4.8,.38,'سپاس از توجه شما',18,GOLD,True); footer(s,18)

# Manual-only slide transitions: no advTm anywhere.
for i,sl in enumerate(prs.slides,1):
    tr=OxmlElement('p:transition'); tr.set('spd','med'); tr.set('advClick','1')
    if i in {3,8,12,16}: child=OxmlElement('p:push'); child.set('dir','l')
    elif i in {2,18}: child=OxmlElement('p:fade')
    elif i in {5,9,13,17}: child=OxmlElement('p:wipe'); child.set('dir','l')
    else: child=OxmlElement('p:split'); child.set('orient','vert'); child.set('dir','out')
    tr.append(child); root=sl._element; pos=1
    for j,e in enumerate(root):
        if e.tag.endswith('clrMapOvr'): pos=j+1
    root.insert(pos,tr)
prs.core_properties.title='چرا به قانون نیاز داریم؟ — نسخه نهایی ۱۸ اسلایدی'
prs.core_properties.subject='ارائه ده دقیقه‌ای کلاس هفتم؛ مدرن، تمیز، دستی و سبک'
prs.core_properties.author='گروه ۶ کلاس ۷/۳'
prs.core_properties.comments='بدون پیشروی خودکار؛ دارای GIF کوچک و صدای کوتاه در نقاط فصل‌بندی'
prs.save(OUT)

# Inject restrained transition sounds on chapter starts and ending.
PNS='http://schemas.openxmlformats.org/presentationml/2006/main'; RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'; RELNS='http://schemas.openxmlformats.org/package/2006/relationships'
tmp=OUT+'.tmp'
with zipfile.ZipFile(OUT) as zin:
    data={n:zin.read(n) for n in zin.namelist()}
for sn,kind in {3:'whoosh',8:'whoosh',12:'chime',16:'whoosh',18:'chime'}.items():
    rel=f'ppt/slides/_rels/slide{sn}.xml.rels'
    rr=ET.fromstring(data[rel]); used={e.get('Id') for e in rr}; k=1
    while f'rIdSfx{k}' in used:k+=1
    rid=f'rIdSfx{k}'
    ET.SubElement(rr,f'{{{RELNS}}}Relationship',{'Id':rid,'Type':'http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio','Target':f'../media/final_{kind}.wav'})
    data[rel]=ET.tostring(rr,encoding='utf-8',xml_declaration=True)
    sx=f'ppt/slides/slide{sn}.xml'; sr=ET.fromstring(data[sx]); tr=sr.find(f'{{{PNS}}}transition')
    sndAc=ET.SubElement(tr,f'{{{PNS}}}sndAc'); st=ET.SubElement(sndAc,f'{{{PNS}}}stSnd'); snd=ET.SubElement(st,f'{{{PNS}}}snd'); snd.set(f'{{{RNS}}}embed',rid); snd.set('name',kind)
    data[sx]=ET.tostring(sr,encoding='utf-8',xml_declaration=True)
ct=data['[Content_Types].xml']
if b'Extension="wav"' not in ct: ct=ct.replace(b'</Types>',b'<Default Extension="wav" ContentType="audio/wav"/></Types>')
data['[Content_Types].xml']=ct
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as zout:
    for n,b in data.items(): zout.writestr(n,b)
    zout.write(f'{V}/whoosh.wav','ppt/media/final_whoosh.wav'); zout.write(f'{V}/chime.wav','ppt/media/final_chime.wav')
os.replace(tmp,OUT)
print(OUT,round(os.path.getsize(OUT)/1048576,2),'MB')
