from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageDraw, ImageEnhance
import os

OUT='output/چرا_به_قانون_نیاز_داریم_نسخه_۸_تیک‌تاک_کنفرانس.pptx'
A='assets_v3'; V='assets_v8'; os.makedirs(V,exist_ok=True); os.makedirs('output',exist_ok=True)
SW,SH=13.333,7.5
BLACK=(3,9,15); NAVY=(7,22,35); WHITE=(246,250,252); ICE=(198,231,245); GOLD=(245,185,66); MINT=(77,210,172); CORAL=(245,103,89); MUTED=(124,151,165)
FONT='Tahoma'; FA=str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')
src={'justice':f'{A}/hero_justice.png','school':f'{A}/hero_school_clean.png','city':f'{A}/hero_city_clean.png','traffic':f'{A}/traffic_3d.png','camp':f'{A}/camp_3d.png','constitution':f'{A}/constitution_3d.png'}

def prep(path,out,side='right'):
    im=Image.open(path).convert('RGB'); w,h=1920,1080; sc=max(w/im.width,h/im.height); im=im.resize((int(im.width*sc),int(im.height*sc)),Image.Resampling.LANCZOS); im=im.crop(((im.width-w)//2,(im.height-h)//2,(im.width+w)//2,(im.height+h)//2)); im=ImageEnhance.Contrast(im).enhance(1.08)
    ov=Image.new('RGBA',(w,h),(0,0,0,0)); d=ImageDraw.Draw(ov)
    for x in range(w):
        t=x/(w-1); a=int(235*((1-t)**1.45 if side=='left' else t**1.45)); d.line((x,0,x,h),fill=(2,7,12,a))
    Image.alpha_composite(im.convert('RGBA'),ov).convert('RGB').save(out,'JPEG',quality=95,subsampling=0)
for n,p in src.items():
    prep(p,f'{V}/{n}_left.jpg','left'); prep(p,f'{V}/{n}_right.jpg','right')

prs=Presentation(); prs.slide_width=Inches(SW); prs.slide_height=Inches(SH); blank=prs.slide_layouts[6]
def C(c): return RGBColor(*c)
def name(sh,n):
    try:
        if sh.shape_type==13: sh._element.nvPicPr.cNvPr.set('name',n)
        else: sh._element.nvSpPr.cNvPr.set('name',n)
    except: pass
    return sh

def rtl(p,center=False):
    p.alignment=PP_ALIGN.CENTER if center else PP_ALIGN.RIGHT; pr=p._p.get_or_add_pPr(); pr.set('rtl','1'); pr.set('algn','ctr' if center else 'r')
def txt(sl,x,y,w,h,s,size=22,col=WHITE,bold=False,center=False,top=False,n=None):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP if top else MSO_ANCHOR.MIDDLE; tf.margin_left=tf.margin_right=Inches(.03); tf.margin_top=tf.margin_bottom=Inches(.02)
    p=tf.paragraphs[0]; p.text=s; rtl(p,center); p.line_spacing=1.05
    for r in p.runs: r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=C(col); r._r.get_or_add_rPr().set('lang','fa-IR')
    if n: name(sh,n)
    return sh

def shadow(sh,alpha=48,blur=62000,dist=26000):
    sp=sh._element.spPr; eff=sp.find(qn('a:effectLst'))
    if eff is None: eff=OxmlElement('a:effectLst'); sp.append(eff)
    o=OxmlElement('a:outerShdw'); o.set('blurRad',str(blur)); o.set('dist',str(dist)); o.set('dir','2700000'); o.set('algn','ctr'); o.set('rotWithShape','0')
    c=OxmlElement('a:srgbClr'); c.set('val','000000'); a=OxmlElement('a:alpha'); a.set('val',str(alpha*1000)); c.append(a); o.append(c); eff.append(o)

def box(sl,kind,x,y,w,h,fill,line=None,trans=0,n=None,shadowed=False):
    sh=sl.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=C(fill); sh.fill.transparency=trans
    if line: sh.line.color.rgb=C(line); sh.line.width=Pt(1.1)
    else: sh.line.fill.background()
    if shadowed: shadow(sh)
    if n: name(sh,n)
    return sh

def glass(sl,x,y,w,h,accent=GOLD,n=None):
    box(sl,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,(5,20,32),(92,121,136),12,n,True)
    box(sl,MSO_SHAPE.ROUNDED_RECTANGLE,x+.03,y+.03,w-.06,.055,(210,231,239),None,82)
    box(sl,MSO_SHAPE.ROUNDED_RECTANGLE,x+w-.055,y+.2,.035,h-.4,accent,None)

def chip(sl,x,y,w,label,accent=GOLD):
    box(sl,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,.34,(11,34,49),accent,4,shadowed=True)
    txt(sl,x+.08,y+.03,w-.16,.25,label,9.5,accent,True,True)

def pic(sl,key,side='right',n='!!Scene'):
    p=sl.shapes.add_picture(f'{V}/{key}_{side}.jpg',0,0,width=prs.slide_width,height=prs.slide_height); name(p,n); return p

def accent_for(n):
    if n<=3: return GOLD
    if n<=9: return MINT
    if n<=13: return CORAL
    if n<=16: return (178,135,245)
    return GOLD

def base(n,label,title,image=None,side='right'):
    s=prs.slides.add_slide(blank); ac=accent_for(n)
    if image: pic(s,image,side)
    else: box(s,MSO_SHAPE.RECTANGLE,0,0,SW,SH,BLACK)
    # Editorial title strip; less repetitive than a full-width card.
    box(s,MSO_SHAPE.RECTANGLE,0,0,SW,1.42,BLACK,None,26,'!!TitleShade')
    dx=.28 if n%2==0 else 0
    chip(s,.68+dx,.29,1.75,label,ac)
    txt(s,.68+dx,.74,11.65-dx,.48,title,29,WHITE,True,False,False,'!!Title')
    box(s,MSO_SHAPE.RECTANGLE,.68+dx,1.3,1.45+.22*(n%3),.028,ac,None,0,'!!TitleRule')
    # TikTok-inspired travelling portal: one shape, new camera position on each slide.
    visual_left = (side=='right') if image else False
    px=(.18+.2*(n%3)) if visual_left else (10.95-.18*(n%3))
    py=2.0+.34*(n%4); ps=1.45+.22*(n%3)
    portal=box(s,MSO_SHAPE.OVAL,px,py,ps,ps,BLACK,ac,100,'!!Portal')
    portal.line.width=Pt(3.2)
    # Kinetic edge rail carries direction without object animations.
    rx=12.94 if n%2==0 else .30
    box(s,MSO_SHAPE.RECTANGLE,rx,1.72,.035,3.92,ac,None,0,'!!KineticRail')
    txt(s,12.05,.32,.48,.35,str(n).zfill(2).translate(FA),11,ac,True,True,'!!SlideNo')
    footer(s,n)
    return s

def footer(sl,n):
    txt(sl,.66,7.12,2.8,.17,'مطالعات اجتماعی • درس ۳',8,MUTED)
    txt(sl,11.78,7.03,.86,.25,str(n).zfill(2).translate(FA)+' / ۱۸',10,GOLD,True)
    # one thin progress line, not decorative clutter
    box(sl,MSO_SHAPE.RECTANGLE,.66,6.88,12.0,.018,(45,66,77))
    box(sl,MSO_SHAPE.RECTANGLE,.66,6.86,12.0*n/18,.04,accent_for(n),n='!!Progress')

def headline(sl,x,y,w,lead,sub,col=GOLD):
    glass(sl,x-.22,y-.2,w+.44,2.32,col)
    # large editorial quotation mark creates a typographic focal point
    txt(sl,x+w-.62,y-.12,.52,.52,'«',31,col,True,True)
    txt(sl,x,y,w-.45,.72,lead,28,col,True)
    box(sl,MSO_SHAPE.RECTANGLE,x,y+.77,1.18,.028,col)
    txt(sl,x,y+.96,w,1.02,sub,17,ICE,False,False,True)

def metric(sl,x,y,num,label,col=GOLD):
    box(sl,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,3.25,.68,(8,31,46),col,3,shadowed=True)
    txt(sl,x+.12,y+.06,.92,.5,num,24,col,True,True)
    box(sl,MSO_SHAPE.RECTANGLE,x+1.13,y+.17,.018,.34,col)
    txt(sl,x+1.32,y+.1,1.68,.4,label,12.5,ICE,True)

# 1 cover — ultra clean
s=prs.slides.add_slide(blank); pic(s,'justice','left')
box(s,MSO_SHAPE.RECTANGLE,.7,.62,.045,5.9,GOLD,n='!!Rail')
txt(s,1.0,.72,5.7,.28,'درس ۳ • مطالعات اجتماعی پایه هفتم',10.5,GOLD,True)
txt(s,.98,1.48,5.8,1.42,'چرا به قانون\nنیاز داریم؟',37,WHITE,True,False,False,'!!Hero')
txt(s,1.0,3.25,5.35,.65,'مسیرِ نظم، امنیت و عدالت',19,ICE)
box(s,MSO_SHAPE.RECTANGLE,1.0,4.12,4.65,.02,(86,112,126))
txt(s,1.0,4.45,5.15,.65,'گروه ۶ • کلاس ۷/۳\nامین کریم‌پور · عزیزی · طاها عابدی‌فر · گله‌دار · محمدرضا فولادی',11.5,WHITE)
footer(s,1)
# 2 book-like chapter menu — roadmap without presenter names
s=base(2,'فهرست مسیر','در این ارائه چه می‌خوانیم؟')
box(s,MSO_SHAPE.RECTANGLE,1.17,1.62,.025,4.55,GOLD,None,0,'!!MenuSpine')
chapters=[('۱','آغاز و مسئله','اسلایدهای ۱ تا ۳',GOLD),('۲','چرا به قانون نیاز داریم؟','اسلایدهای ۴ تا ۹',MINT),('۳','قانون در زندگی روزمره','اسلایدهای ۱۰ تا ۱۳',CORAL),('۴','حق، تکلیف و چارچوب کشور','اسلایدهای ۱۴ تا ۱۶',(178,135,245)),('۵','مرور و جمع‌بندی','اسلایدهای ۱۷ و ۱۸',GOLD)]
for i,(num,title,rng,col) in enumerate(chapters):
    y=1.55+i*.92; x=.92+i*.34
    txt(s,x,y,.55,.5,num,21,col,True,True)
    box(s,MSO_SHAPE.RECTANGLE,x+.78,y+.5,8.7-i*.22,.018,col)
    txt(s,x+.86,y-.01,6.7,.42,title,17,WHITE,True)
    txt(s,10.2,y+.05,1.95,.3,rng,10.5,MUTED,True,True)
# 3 social life
s=base(3,'زندگی اجتماعی','هر ارتباط، یک حق و یک مسئولیت می‌سازد','school','right')
headline(s,7.1,1.7,5.35,'حقِ من ↔ مسئولیتِ من','در خانواده، مدرسه و محله از امکانات و احترام برخورداریم؛ در مقابل باید حق دیگران را رعایت کنیم.',MINT)
metric(s,7.12,4.55,'۰۱','حق آموزش',MINT); metric(s,9.35,5.25,'۰۲','تکلیف رعایت نظم',GOLD)
# 4 boundary freedom
s=base(4,'اصل طلایی','مرز آزادی من، حق دیگران است')
txt(s,.9,4.9,11.6,1.25,'آزادی',64,(30,45,55),True,True,n='!!GhostWord')
glass(s,.72,1.52,11.9,1.22,GOLD)
txt(s,.9,1.72,11.55,.72,'«آزادی بدون مسئولیت، به بی‌نظمی تبدیل می‌شود.»',27,GOLD,True,True)
box(s,MSO_SHAPE.RECTANGLE,2.0,3.05,9.3,.02,(64,91,105))
for i,(h,b,c) in enumerate([('حق من','انتخاب و استفاده از امکانات',MINT),('قانون','مرز مشترک رفتار',GOLD),('حق دیگران','آرامش، امنیت و دارایی',CORAL)]):
    x=1.12+i*4.05; txt(s,x,3.62,3.35,.38,h,18,c,True,True); txt(s,x,4.25,3.35,.58,b,13.5,ICE,False,True)
# 5 compare
s=base(5,'تفکیک','مقررات و قانون؛ تفاوت در گستره و اجرا')
box(s,MSO_SHAPE.RECTANGLE,6.66,1.55,.015,4.72,GOLD)
headline(s,.92,1.72,5.15,'مقررات','برای محیط محدود؛ مثل خانه، کلاس و اردو. معمولاً با توافق گروه یا مدیریت همان محیط اجرا می‌شود.',MINT)
headline(s,7.2,1.72,5.15,'قانون','قاعده‌ای رسمی و لازم‌الاجرا برای جامعه؛ با نظارت نهادهای رسمی و دادگاه.',GOLD)
# 6 four reasons, simple numbered list
s=base(6,'هستهٔ درس','چهار دلیل اصلی نیاز به قانون','justice','right')
glass(s,6.82,1.48,5.55,4.62,GOLD)
chip(s,10.55,1.7,1.45,'چهار ستون')
items=[('۱','حفظ حقوق افراد'),('۲','برقراری نظم و امنیت'),('۳','رشد مسئولیت‌پذیری'),('۴','حل عادلانهٔ اختلاف')]
for i,(n,t) in enumerate(items):
    y=1.65+i*1.05; txt(s,7.15,y,.55,.45,n,17,GOLD,True,True); txt(s,7.92,y,4.15,.45,t,16,WHITE,True); box(s,MSO_SHAPE.RECTANGLE,7.92,y+.62,3.8,.015,(56,79,91))
# 7 rights protection
s=base(7,'دلیل اول','قانون، سپر افراد در برابر زورگویی است','justice','left')
txt(s,7.15,4.55,5.2,1.15,'حــق',58,(48,59,63),True,True,n='!!GhostWord')
headline(s,.88,1.75,5.15,'قدرت، جای حق را نمی‌گیرد.','قانون به همه—حتی افراد ضعیف‌تر—امکان دفاع، شکایت و رسیدگی عادلانه می‌دهد.',GOLD)
# 8 traffic
s=base(8,'دلیل دوم','چراغ قرمز؛ قرارداد حفظ جان','traffic','right')
txt(s,.55,4.78,5.7,1.1,'جــان',57,(35,48,53),True,True,n='!!GhostWord')
headline(s,7.1,1.72,5.25,'چند ثانیه عجله، چند حق از دست‌رفته','عبور غیرمجاز فقط راننده را درگیر نمی‌کند؛ جان، زمان و آرامش دیگران هم تهدید می‌شود.',CORAL)
metric(s,7.12,4.55,'جان','امنیت',CORAL); metric(s,9.6,5.18,'زمان','نظم',GOLD)
# 9 responsibility + court
s=base(9,'دلیل سوم و چهارم','مسئولیت‌پذیری و حل اختلاف')
headline(s,.92,1.65,5.2,'مسئولیت‌پذیری','قانون به ما یاد می‌دهد پیامد رفتار خود را ببینیم و حتی بدون نظارت، انتخاب درست داشته باشیم.',MINT)
headline(s,7.18,1.65,5.2,'حل اختلاف','دادگاه، دلیل و قانون را جایگزین دعوا، زور و انتقام می‌کند.',GOLD)
# 10 public property
s=base(10,'قانون در شهر','مال عمومی، مال همه است','city','left')
headline(s,.88,1.7,5.25,'پارک • اتوبوس • کتابخانه','پاکیزگی، مراقبت از تجهیزات و امانت‌داری باعث می‌شود امکانات برای همه و نسل بعد باقی بماند.',MINT)
# 11 camp intro
s=base(11,'کارگاه صفحهٔ ۱۵','اردو؛ قانون خوب باید قابل اجرا باشد','camp','right')
headline(s,7.08,1.7,5.22,'سه معیار ساده','قانون خوب باید روشن، عادلانه و قابل اجرا باشد؛ نه فقط سخت‌گیرانه.',GOLD)
# 12 camp rules
s=base(12,'مقررات اردو','پنج موقعیت، پنج تصمیم روشن','camp','right')
glass(s,6.78,1.38,5.62,4.95,MINT)
chip(s,10.78,1.58,1.25,'چک‌لیست',MINT)
rules=[('وسایل','جای مشخص و تحویل سالم'),('استراحت','خاموشی و سکوت'),('آتش','محل مجاز و حضور مسئول'),('غذا','صف و سهم برابر'),('گروه','نقش روشن و همکاری')]
for i,(h,b) in enumerate(rules):
    y=1.45+i*.92; txt(s,7.05,y,1.3,.34,h,13,GOLD,True); txt(s,8.55,y,3.45,.34,b,13.5,WHITE); box(s,MSO_SHAPE.RECTANGLE,7.05,y+.56,4.95,.012,(51,75,88))
# 13 ethics
s=base(13,'قانون و اخلاق','حداقلِ لازم؛ بهترینِ ممکن')
headline(s,.92,1.62,5.2,'قانون','رفتار لازم و الزام‌آور؛ مثل ایست پشت چراغ قرمز.',GOLD)
headline(s,7.18,1.62,5.2,'اخلاق','رفتار انسانی و برتر؛ مثل کمک داوطلبانه به فرد نابینا.',MINT)
txt(s,2.65,5.35,8.05,.44,'جامعهٔ خوب، هم قانون‌مدار است و هم اخلاق‌مدار.',17,WHITE,True,True)
# 14 constitution
s=base(14,'دانستنی','قانون اساسی؛ مادر قوانین','constitution','left')
txt(s,7.18,4.72,5.1,1.1,'چارچوب',50,(45,45,62),True,True,n='!!GhostWord')
headline(s,.88,1.62,5.25,'چارچوب اصلی کشور','قانون اساسی حقوق عمومی و ساختار نهادها را مشخص می‌کند. قوانین دیگر نباید با آن مغایر باشند.',GOLD)
txt(s,.9,4.75,4.9,.38,'همه‌پرسی: ۱۲ آذر ۱۳۵۸',14,MINT,True,True)
# 15 right duty
s=base(15,'تعادل','حق و تکلیف؛ دو کفهٔ یک ترازو','justice','right')
txt(s,.55,4.75,5.65,1.1,'تعادل',52,(38,47,56),True,True,n='!!GhostWord')
headline(s,7.1,1.52,5.2,'حق','آموزش، امنیت، احترام و استفاده از امکانات',MINT)
box(s,MSO_SHAPE.RECTANGLE,7.1,3.38,4.6,.02,(78,101,113))
headline(s,7.1,3.65,5.2,'تکلیف','رعایت نظم، حفظ اموال و احترام به دیگران',CORAL)
# 16 decision framework
s=base(16,'تصمیم مسئولانه','پیش از عمل، سه سؤال بپرس')
glass(s,1.62,1.48,10.1,4.45,GOLD)
chip(s,9.75,1.7,1.48,'مکث کن')
questions=[('۱','آیا کسی آسیب می‌بیند؟'),('۲','اگر همه انجام دهند چه می‌شود؟'),('۳','آیا پیامدش را می‌پذیرم؟')]
for i,(n,q) in enumerate(questions):
    y=1.7+i*1.35; txt(s,2.0,y,.65,.48,n,18,GOLD,True,True); txt(s,2.95,y,8.2,.48,q,19,WHITE,True); box(s,MSO_SHAPE.RECTANGLE,2.95,y+.7,7.2,.015,(59,83,96))
# 17 review — Karim easy
s=base(17,'مرور سریع','سه نکته‌ای که باید بماند')
glass(s,1.08,1.46,11.15,4.38,MINT)
chip(s,10.3,1.68,1.5,'خلاصهٔ درس',MINT)
for i,(h,c) in enumerate([('قانون از حق‌ها محافظت می‌کند.',GOLD),('آزادی با مسئولیت کامل می‌شود.',MINT),('حق و تکلیف از هم جدا نیستند.',CORAL)]):
    y=1.65+i*1.38; txt(s,1.45,y,.6,.48,str(i+1).translate(FA),18,c,True,True); txt(s,2.35,y,9.4,.48,h,21,WHITE,True)
# 18 ending — Karim easy
s=prs.slides.add_slide(blank); pic(s,'justice','left'); box(s,MSO_SHAPE.RECTANGLE,.7,.62,.045,5.9,GOLD,n='!!Rail')
txt(s,1.0,.78,4.6,.25,'جمع‌بندی نهایی',10.5,GOLD,True)
txt(s,.98,1.58,5.85,1.52,'قانون یعنی\nاحترام به حقِ خود و دیگران.',31,WHITE,True,False,False,'!!Hero')
txt(s,1.0,3.65,5.25,.7,'جامعهٔ بهتر، از انتخاب مسئولانهٔ ما آغاز می‌شود.',16,ICE)
txt(s,1.0,5.05,4.8,.38,'سپاس از توجه شما',18,GOLD,True)
footer(s,18)

# TikTok/conference motion system: one visual language, no in-slide animations.
P159='http://schemas.microsoft.com/office/powerpoint/2015/09/main'
for i,s in enumerate(prs.slides,1):
    tr=OxmlElement('p:transition'); tr.set('spd','med'); tr.set('advClick','1')
    if i in {4,10,14,17}:
        z=OxmlElement('p:zoom'); z.set('dir','in'); tr.append(z)
    elif i==1:
        f=OxmlElement('p:fade'); tr.append(f)
    else:
        m=etree.Element('{%s}morph'%P159,nsmap={'p159':P159}); m.set('option','byObject'); tr.append(m)
    root=s._element; pos=1
    for j,e in enumerate(root):
        if e.tag.endswith('clrMapOvr'): pos=j+1
    root.insert(pos,tr)
prs.core_properties.title='چرا به قانون نیاز داریم؟ — نسخه ۸ تیک‌تاک کنفرانس'
prs.core_properties.subject='ارائهٔ هیبرید ادیتوریال، سه‌بعدی و وایرال با حرکت دوربین'
prs.core_properties.author='گروه ۶ از کلاس ۷/۳'
prs.core_properties.comments='نسخه ۸: تکنیک‌های منتخب آموزش‌های TikTok و کنفرانس؛ پورتال Morph، زوم فصل‌ها و تایپوگرافی ادیتوریال؛ بدون انیمیشن داخلی'
prs.save(OUT); print(OUT,round(os.path.getsize(OUT)/1048576,2),'MB')
