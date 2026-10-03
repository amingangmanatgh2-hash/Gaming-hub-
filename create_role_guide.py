from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import arabic_reshaper
from bidi.algorithm import get_display
import re, os

SRC='راهنمای_کامل_تقسیم_نقش_نسخه۵.md'
DOCX='output/راهنمای_تقسیم_نقش_و_متن_ارائه_نسخه۵.docx'
PDF='output/راهنمای_تقسیم_نقش_و_متن_ارائه_نسخه۵.pdf'
text=open(SRC,encoding='utf-8').read()
os.makedirs('output',exist_ok=True)

# ---------- DOCX ----------
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
styles=doc.styles
for name,size,bold,color in [('Normal',11,False,'183047'),('Title',25,True,'0B2439'),('Heading 1',20,True,'0B2439'),('Heading 2',15,True,'D39B2E')]:
    st=styles[name]; st.font.name='Tahoma'; st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor.from_string(color); st._element.rPr.rFonts.set(qn('w:cs'),'Tahoma'); st._element.rPr.rFonts.set(qn('w:eastAsia'),'Tahoma')

def rtlp(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    pPr=p._p.get_or_add_pPr(); bidi=OxmlElement('w:bidi'); pPr.append(bidi)
    p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.12

def addp(s,style=None,bold=False,color=None):
    p=doc.add_paragraph(style=style); rtlp(p); r=p.add_run(s); r.font.name='Tahoma'; r._element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Tahoma'); r.bold=bold
    if color: r.font.color.rgb=RGBColor.from_string(color)
    return p

lines=text.splitlines(); i=0
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('|') and i+1<len(lines) and set(lines[i+1].replace('|','').replace('-','').replace(':','').strip())==set():
        rows=[]; i+=2
        headers=[x.strip() for x in line.strip('|').split('|')]
        while i<len(lines) and lines[i].strip().startswith('|'):
            rows.append([x.strip() for x in lines[i].strip().strip('|').split('|')]); i+=1
        table=doc.add_table(rows=1,cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Light Shading Accent 1'
        for j,h in enumerate(headers): table.rows[0].cells[j].text=h
        for row in rows:
            cells=table.add_row().cells
            for j,v in enumerate(row): cells[j].text=v
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcPr=cell._tc.get_or_add_tcPr(); td=OxmlElement('w:textDirection'); td.set(qn('w:val'),'rtl'); tcPr.append(td)
                for p in cell.paragraphs:
                    rtlp(p)
                    for r in p.runs: r.font.name='Tahoma'; r.font.size=Pt(9); r._element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Tahoma')
        continue
    if line.startswith('# '): addp(line[2:],'Title');
    elif line.startswith('## '):
        if line.startswith('## اسلاید') and len(doc.paragraphs)>3: pass
        addp(line[3:],'Heading 1')
    elif line.startswith('# '): addp(line[2:],'Heading 1')
    elif line.startswith('---'): doc.add_paragraph('')
    elif line.startswith('- '): addp('• '+line[2:])
    elif re.match(r'^\d+\. ',line): addp(line)
    else:
        clean=line.replace('**','').replace('  ',' ')
        is_label=clean.endswith(':') or clean in ('چگونه بگوید:','تحویل به نفر بعد:')
        addp(clean,bold=is_label,color='D39B2E' if is_label else None)
    i+=1
# header/footer
for section in doc.sections:
    hp=section.header.paragraphs[0]; hp.text='گروه ۶ • کلاس ۷/۳ | راهنمای اجرای نسخهٔ ۵'; rtlp(hp)
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)
doc.save(DOCX)

# ---------- PDF as crisp paged images ----------
W,H=1440,2037; M=95; MAXW=W-2*M
reg='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'; boldf='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
F={k:ImageFont.truetype(boldf if b else reg,s) for k,s,b in [('title',42,1),('h1',33,1),('h2',27,1),('body',22,0),('bold',22,1),('small',18,0)]}

def visual(s): return get_display(arabic_reshaper.reshape(s))
def cleanmd(s): return s.replace('**','').replace('`','').replace('  ',' ').strip()
def wrap(s,font,maxw=MAXW):
    words=s.split(); out=[]; cur=''
    for word in words:
        test=(cur+' '+word).strip(); width=ImageDraw.Draw(Image.new('RGB',(1,1))).textbbox((0,0),visual(test),font=font)[2]
        if width<=maxw: cur=test
        else:
            if cur: out.append(cur)
            cur=word
    if cur: out.append(cur)
    return out or ['']

pages=[]; im=None; d=None; y=0

def new_page():
    global im,d,y
    if im is not None: pages.append(im)
    im=Image.new('RGB',(W,H),(242,246,246)); d=ImageDraw.Draw(im); y=88
    d.rounded_rectangle((55,45,W-55,H-45),28,fill=(250,252,252),outline=(205,218,220),width=3)
    d.rectangle((W-72,45,W-55,H-45),fill=(255,198,78))

def draw_lines(s,font,color=(20,45,65),gap=12,indent=0):
    global y
    for line in wrap(s,font,MAXW-indent):
        if y+font.size+20>H-95: new_page()
        v=visual(line); box=d.textbbox((0,0),v,font=font); d.text((W-M-indent,y),v,font=font,fill=color,anchor='ra'); y+=font.size+gap

def rule():
    global y
    d.line((M,y,W-M,y),fill=(213,155,46),width=3); y+=20

new_page()
for raw in lines:
    line=raw.strip()
    if not line: y+=8; continue
    if line=='---': rule(); continue
    if line.startswith('# '):
        if y>200: new_page()
        draw_lines(cleanmd(line[2:]),F['title'],(8,38,61),18); rule()
    elif line.startswith('## '):
        # Start each member section on a fresh page
        heading=cleanmd(line[3:])
        if re.match(r'^[۱-۵]\)',heading) and y>180: new_page()
        y+=18; draw_lines(heading,F['h1'],(8,52,78),15)
    elif line.startswith('### '): y+=10; draw_lines(cleanmd(line[4:]),F['h2'],(203,139,35),12)
    elif line.startswith('|'):
        # table rendered as readable text rows
        if set(line.replace('|','').replace('-','').replace(':','').strip())==set(): continue
        vals=[x.strip() for x in line.strip('|').split('|')]
        draw_lines(' • '.join(vals),F['small'],(38,73,92),8,20)
    elif line.startswith('- '): draw_lines('• '+cleanmd(line[2:]),F['body'],(23,52,70),11,18)
    elif re.match(r'^\d+\. ',line): draw_lines(cleanmd(line),F['body'],(23,52,70),11,18)
    else:
        c=cleanmd(line); lab=c.endswith(':') or c in ('چگونه بگوید:','تحویل به نفر بعد:')
        draw_lines(c,F['bold'] if lab else F['body'],(203,139,35) if lab else (23,52,70),11)
if im is not None: pages.append(im)
for i,p in enumerate(pages,1):
    dd=ImageDraw.Draw(p); dd.text((W//2,H-70),str(i).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')),font=F['small'],fill=(115,135,145),anchor='mm')
pages[0].save(PDF,'PDF',resolution=150,save_all=True,append_images=pages[1:])
print(DOCX,PDF,len(pages))
