from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFilter
import os, math, random, wave, struct, zipfile, shutil, tempfile

OUT='output/چرا_به_قانون_نیاز_داریم_گروه۶.pptx'
AS='assets'
os.makedirs('output',exist_ok=True); os.makedirs(AS,exist_ok=True)
W,H=1600,900
NAVY=(4,16,30); NAVY2=(7,27,49); ICE=(200,235,255); GOLD=(255,201,92); MINT=(88,224,188); CORAL=(255,119,105); WHITE=(239,248,255)
FONT='Vazirmatn'

# ---------- art ----------
def bg(path, seed):
    random.seed(seed)
    im=Image.new('RGB',(W,H),NAVY); px=im.load()
    for y in range(H):
        t=y/H
        for x in range(W):
            g=int(10*(x/W)+7*t)
            px[x,y]=(4+g//4,16+g//2,30+g)
    glow=Image.new('RGBA',(W,H),(0,0,0,0)); d=ImageDraw.Draw(glow)
    for cx,cy,c in [(1300,130,(40,150,255)),(180,760,(24,190,160))]:
        for r in range(320,10,-18):
            a=max(0,int(2.6*(330-r)))//10
            d.ellipse((cx-r,cy-r,cx+r,cy+r),fill=(*c,min(10,a)))
    glow=glow.filter(ImageFilter.GaussianBlur(35)); im=Image.alpha_composite(im.convert('RGBA'),glow)
    d=ImageDraw.Draw(im)
    for i in range(55):
        x=random.randrange(W); y=random.randrange(H); r=random.choice([1,1,2]); d.ellipse((x-r,y-r,x+r,y+r),fill=(180,225,255,random.randrange(25,90)))
    d.line([(0,820),(1600,590)],fill=(70,160,210,30),width=2)
    d.line([(600,900),(1600,520)],fill=(255,201,92,25),width=2)
    im.convert('RGB').save(path,quality=92)

def layer(): return Image.new('RGBA',(600,600),(0,0,0,0))
def glow(draw,xy,color,width=18):
    for w in range(width,0,-3): draw.line(xy,fill=(*color,max(15,90-w*3)),width=w)

def icon_scale(path):
    im=layer(); sh=Image.new('RGBA',im.size,(0,0,0,0)); d=ImageDraw.Draw(sh)
    d.ellipse((115,485,505,555),fill=(0,0,0,130)); sh=sh.filter(ImageFilter.GaussianBlur(20)); im=Image.alpha_composite(im,sh); d=ImageDraw.Draw(im)
    # pedestal and beam
    d.rounded_rectangle((245,430,355,505),20,fill=(160,95,18),outline=GOLD,width=8)
    d.polygon([(300,80),(258,430),(342,430)],fill=(175,105,20),outline=GOLD)
    d.ellipse((265,52,335,122),fill=(255,225,130),outline=GOLD,width=8)
    d.rounded_rectangle((105,145,495,180),17,fill=(211,142,35),outline=(255,231,150),width=7)
    for x in (165,435):
        d.line((x,175,x-58,350),fill=GOLD,width=8); d.line((x,175,x+58,350),fill=GOLD,width=8)
        d.arc((x-105,305,x+105,425),0,180,fill=(255,225,130),width=12)
        d.line((x-104,365,x+104,365),fill=GOLD,width=6)
    im.save(path)

def icon_check(path):
    im=layer(); d=ImageDraw.Draw(im)
    d.rounded_rectangle((80,85,340,510),35,fill=(18,55,84,245),outline=ICE,width=8)
    for y in (180,270,360):
        d.rounded_rectangle((120,y,165,y+45),8,outline=GOLD,width=6); d.line((128,y+23,142,y+37,158,y+10),fill=MINT,width=7); d.line((190,y+20,300,y+20),fill=ICE,width=7)
    d.polygon([(420,130),(520,170),(500,350),(420,430),(340,350),(320,170)],fill=(28,115,150),outline=(120,235,255),width=8)
    d.line((372,280,405,315,470,225),fill=GOLD,width=16)
    d.polygon([(390,470),(510,120),(548,138),(430,490)],fill=GOLD,outline=(255,235,170)); d.polygon([(390,470),(430,490),(380,520)],fill=(255,170,100))
    im.save(path)

def icon_network(path):
    im=layer(); d=ImageDraw.Draw(im)
    pts=[(300,130),(145,270),(455,270),(215,450),(385,450)]
    for a,b in [(0,1),(0,2),(1,2),(1,3),(2,4),(3,4),(1,4),(2,3)]: d.line((pts[a],pts[b]),fill=(90,190,235,150),width=7)
    colors=[GOLD,MINT,CORAL,ICE,GOLD]
    for (x,y),c in zip(pts,colors):
        d.ellipse((x-44,y-44,x+44,y+44),fill=(*c,245),outline=WHITE,width=5); d.ellipse((x-16,y-25,x+16,y+7),fill=NAVY); d.arc((x-27,y,x+27,y+43),180,360,fill=NAVY,width=12)
    im.save(path)

def icon_class(path):
    im=layer(); d=ImageDraw.Draw(im)
    d.polygon([(80,210),(485,150),(535,420),(130,485)],fill=(17,58,84),outline=ICE)
    d.rectangle((145,195,445,340),fill=(10,85,88),outline=(140,255,220),width=8)
    for x,y,c in [(115,110,GOLD),(380,75,CORAL),(430,420,MINT)]:
        d.rounded_rectangle((x,y,x+130,y+85),15,fill=c,outline=WHITE,width=5); d.line((x+20,y+23,x+110,y+23),fill=NAVY,width=5); d.line((x+20,y+48,x+90,y+48),fill=NAVY,width=5)
    for x in (170,320): d.polygon([(x,370),(x+110,350),(x+135,410),(x+20,430)],fill=(128,76,35),outline=GOLD)
    im.save(path)

def icon_traffic(path):
    im=layer(); d=ImageDraw.Draw(im)
    d.polygon([(40,520),(225,290),(565,470),(430,560)],fill=(35,45,55));
    for i in range(5): d.polygon([(180+i*55,405-i*9),(204+i*55,392-i*9),(245+i*55,410-i*9),(220+i*55,425-i*9)],fill=WHITE)
    d.rounded_rectangle((150,45,300,350),34,fill=(19,35,47),outline=ICE,width=8); d.rectangle((210,350,240,510),fill=(110,120,130))
    for y,c,on in [(115,CORAL,True),(200,GOLD,False),(285,MINT,False)]:
        d.ellipse((183,y-35,267,y+49),fill=c if on else tuple(v//4 for v in c),outline=WHITE,width=4)
    im.save(path)

def icon_shield(path):
    im=layer(); d=ImageDraw.Draw(im)
    d.polygon([(300,65),(500,145),(470,385),(300,535),(130,385),(100,145)],fill=(18,90,130),outline=ICE,width=13)
    d.polygon([(300,105),(447,165),(425,360),(300,480),(300,105)],fill=(35,145,172))
    d.line((195,290,270,365,420,200),fill=GOLD,width=28)
    im.save(path)

def icon_camp(path):
    im=layer(); d=ImageDraw.Draw(im)
    d.ellipse((55,445,550,550),fill=(0,0,0,100))
    for x,y,c in [(80,250,CORAL),(265,190,MINT),(400,300,GOLD)]:
        d.polygon([(x,y+190),(x+100,y),(x+210,y+190)],fill=c,outline=WHITE); d.polygon([(x+100,y),(x+125,y+190),(x+210,y+190)],fill=tuple(max(0,v-60) for v in c)); d.polygon([(x+92,y+190),(x+110,y+125),(x+130,y+190)],fill=NAVY)
    d.line((460,180,460,335),fill=(130,85,40),width=12); d.ellipse((430,105,490,190),fill=(255,170,70)); d.ellipse((442,120,478,175),fill=(255,225,100))
    im.save(path)

def icon_ethics(path):
    im=layer(); d=ImageDraw.Draw(im)
    d.polygon([(300,475),(115,300),(110,195),(170,125),(260,140),(300,205),(340,140),(430,125),(490,195),(485,300)],fill=CORAL,outline=(255,190,180),width=8)
    d.rounded_rectangle((175,345,425,430),35,fill=(245,190,125),outline=WHITE,width=6)
    d.line((210,365,300,300,390,365),fill=GOLD,width=12); d.line((300,210,300,345),fill=GOLD,width=12); d.line((230,250,370,250),fill=GOLD,width=12)
    d.arc((185,235,275,320),0,180,fill=GOLD,width=8); d.arc((325,235,415,320),0,180,fill=GOLD,width=8)
    im.save(path)

def icon_doc(path):
    im=layer(); d=ImageDraw.Draw(im)
    d.rounded_rectangle((95,65,475,515),30,fill=(235,224,190),outline=GOLD,width=10)
    d.polygon([(410,65),(475,130),(410,130)],fill=(190,172,130))
    for y,w in [(170,250),(220,290),(270,220),(320,275)]: d.rounded_rectangle((145,y,145+w,y+12),6,fill=(65,80,90))
    d.ellipse((325,355,455,485),fill=(142,35,35),outline=GOLD,width=8); d.ellipse((355,385,425,455),outline=GOLD,width=5)
    im.save(path)

arts={'scale':icon_scale,'check':icon_check,'network':icon_network,'class':icon_class,'traffic':icon_traffic,'shield':icon_shield,'camp':icon_camp,'ethics':icon_ethics,'doc':icon_doc}
for i in range(1,17): bg(f'{AS}/bg{i}.jpg',i)
for n,f in arts.items(): f(f'{AS}/{n}.png')

# ---------- PPT helpers ----------
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]

def set_rtl(p):
    p.alignment=PP_ALIGN.RIGHT
    pPr=p._p.get_or_add_pPr(); pPr.set('rtl','1'); pPr.set('algn','r')

def textbox(slide,x,y,w,h,text,size=18,color=WHITE,bold=False,align='right',margin=.08,valign=MSO_ANCHOR.MIDDLE):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(.03); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.text=text; set_rtl(p); p.alignment={'right':PP_ALIGN.RIGHT,'center':PP_ALIGN.CENTER,'left':PP_ALIGN.LEFT}[align]
    for r in p.runs:
        r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor(*color)
        r._r.get_or_add_rPr().set('lang','fa-IR')
    return sh

def shadow(shape,blur='63500',dist='38100',dirn='2700000',alpha=55):
    spPr=shape._element.spPr; eff=spPr.find(qn('a:effectLst'))
    if eff is None: eff=OxmlElement('a:effectLst'); spPr.append(eff)
    outer=OxmlElement('a:outerShdw'); outer.set('blurRad',blur); outer.set('dist',dist); outer.set('dir',dirn); outer.set('algn','ctr'); outer.set('rotWithShape','0')
    srgb=OxmlElement('a:srgbClr'); srgb.set('val','000000'); al=OxmlElement('a:alpha'); al.set('val',str(alpha*1000)); srgb.append(al); outer.append(srgb); eff.append(outer)

def card(slide,x,y,w,h,fill=(13,42,68),line=(60,140,180),radius=True,accent=None):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(*fill); sh.fill.transparency=8
    sh.line.color.rgb=RGBColor(*line); sh.line.transparency=35; sh.line.width=Pt(1.2); shadow(sh)
    if accent:
        a=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+w-.07),Inches(y+.15),Inches(.055),Inches(h-.3)); a.fill.solid(); a.fill.fore_color.rgb=RGBColor(*accent); a.line.fill.background()
    return sh

def title(slide,num,kicker,ttl):
    textbox(slide,.65,.25,11.9,.35,kicker,10,GOLD,True)
    textbox(slide,.65,.58,11.9,.58,ttl,27,WHITE,True)
    ln=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(10.9),Inches(1.2),Inches(1.75),Inches(.035)); ln.fill.solid(); ln.fill.fore_color.rgb=RGBColor(*GOLD); ln.line.fill.background()

def footer(slide,num):
    textbox(slide,.55,7.08,2.0,.2,'مطالعات اجتماعی • درس ۳',9,(126,176,205),False,'left')
    n=f'{str(num).zfill(2)} / ۱۶'.translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))
    textbox(slide,11.65,7.02,1.05,.3,n,11,GOLD,True,'right')

def base(num,kicker,ttl):
    s=prs.slides.add_slide(blank); s.shapes.add_picture(f'{AS}/bg{num}.jpg',0,0,width=prs.slide_width,height=prs.slide_height); title(s,num,kicker,ttl); footer(s,num); return s

def body_card(slide,x,y,w,h,head,body,accent=GOLD,hs=16,bs=13):
    card(slide,x,y,w,h,accent=accent); textbox(slide,x+.22,y+.16,w-.44,.35,head,hs,accent,True); textbox(slide,x+.22,y+.55,w-.44,h-.68,body,bs,ICE,False,'right',.04,MSO_ANCHOR.TOP)

def add_art(slide,name,x,y,w,h):
    pic=slide.shapes.add_picture(f'{AS}/{name}.png',Inches(x),Inches(y),width=Inches(w),height=Inches(h)); shadow(pic,'90000','45000','2700000',65); return pic

# 1
s=prs.slides.add_slide(blank); s.shapes.add_picture(f'{AS}/bg1.jpg',0,0,width=prs.slide_width,height=prs.slide_height)
textbox(s,.7,.45,5.4,.35,'درس ۳ • مطالعات اجتماعی پایه هفتم',12,GOLD,True)
textbox(s,.7,1.05,7.3,1.2,'چرا به قانون نیاز داریم؟',34,WHITE,True)
textbox(s,.72,2.22,6.3,.7,'قانون، مسیر رسیدن به نظم، امنیت و احترام به حقوق دیگران است.',19,ICE,False)
card(s,.72,3.15,6.15,1.6,accent=MINT)
textbox(s,1.0,3.38,5.5,.35,'ارائه‌شده توسط گروه ۶ از کلاس ۷/۳',16,GOLD,True)
textbox(s,1.0,3.92,5.55,.75,'امین کریم‌پور · عزیزی · طاها عابدی‌فر\nگله‌دار · محمدرضا فولادی',14,WHITE,False)
add_art(s,'scale',7.45,.75,5.1,5.65)
textbox(s,.75,6.5,5.8,.32,'یک ارائهٔ سینمایی دربارهٔ مرز حق، مسئولیت و عدالت',11,MINT,True)
footer(s,1)
#2
s=base(2,'نقشهٔ راه ارائه','از «حق» تا «قانون»؛ مسیر یادگیری ما')
items=[('۰۱','از حق تا قانون','چرا زندگی اجتماعی، مرز و قاعده می‌خواهد؟'),('۰۲','مقررات و نظم','تفاوت قواعد گروهی با قانون عمومی'),('۰۳','امنیت و عدالت','قانون چگونه از افراد و جامعه محافظت می‌کند؟'),('۰۴','حق و تکلیف','دو کفهٔ جدانشدنی زندگی شهروندی')]
for i,(n,h,b) in enumerate(items):
    y=1.55+i*1.18; body_card(s,.65,y,7.1,1.0,h,b,[GOLD,MINT,CORAL,ICE][i],15,11); textbox(s,7.03,y+.22,.48,.4,n,15,[GOLD,MINT,CORAL,ICE][i],True,'center')
add_art(s,'check',8.0,1.35,4.5,4.65)
card(s,8.35,5.82,3.75,.55,fill=(29,61,80),line=GOLD); textbox(s,8.55,5.93,3.35,.28,'نمای بازطراحی‌شده از محورهای صفحهٔ ۱۳',10,ICE,False,'center')
#3
s=base(3,'مفهوم پایه','ورود به اجتماع و شکل‌گیری حقوق متقابل')
body_card(s,.65,1.48,7.35,1.35,'زندگی اجتماعی از کجا آغاز می‌شود؟','ما در خانواده، مدرسه و محله با دیگران رابطه داریم. همین رابطه‌ها سبب می‌شود هر فرد در برابر دیگری «حق» داشته باشد و هم‌زمان مسئول رعایت حق او باشد.',GOLD,16,13)
body_card(s,.65,3.03,7.35,1.55,'اصل طلایی مرز حق','هیچ‌کس در استفاده از حق خود آزادی مطلق و بی‌قیدوشرط ندارد. مرز حق من، جایی است که قانون و حق دیگران آغاز می‌شود؛ پس حق، بدون مسئولیت کامل نیست.',CORAL,16,14)
body_card(s,.65,4.78,7.35,1.35,'مثال کاربردی','دانش‌آموز حق استفاده از حیاط مدرسه را دارد؛ اما نمی‌تواند با بازی خطرناک، امنیت دیگران را تهدید کند. حق بازی + تکلیف مراقبت از دیگران = همزیستی سالم.',MINT,15,12)
add_art(s,'network',8.05,1.4,4.6,4.9)
textbox(s,8.55,6.05,3.55,.34,'ارتباط بیشتر ← حقوق و مسئولیت‌های بیشتر',11,GOLD,True,'center')
#4
s=base(4,'تفکیک مفهومی','مقررات یا قانون؟ شبیه‌اند، اما یکسان نیستند')
body_card(s,.65,1.55,5.45,3.35,'مقررات','• قواعدی برای ادارهٔ یک گروه یا محیط محدود\n• می‌تواند در خانه، کلاس، مدرسه یا اردو وضع شود\n• نمونه: ساعت خاموشی اردو، نوبت نظافت کلاس\n• ضمانت اجرا معمولاً در همان گروه تعریف می‌شود\n\nنتیجه: مقررات، نظمِ نزدیک و روزمره را می‌سازد.',MINT,19,14)
body_card(s,6.2,1.55,5.45,3.35,'قانون','• قاعده‌ای رسمی، عمومی و لازم‌الاجرا\n• به‌وسیلهٔ مرجع قانون‌گذاری تصویب می‌شود\n• برای همهٔ افراد جامعه اعتبار دارد\n• دادگاه و نهادهای رسمی بر اجرای آن نظارت می‌کنند\n\nنتیجه: قانون، نظم و عدالت عمومی را تضمین می‌کند.',GOLD,19,14)
add_art(s,'class',4.65,4.85,4.25,2.0)
textbox(s,.85,5.48,3.6,.56,'وجه مشترک: هر دو برای پیشگیری از آشفتگی و رعایت حقوق دیگران‌اند.',11,ICE,True)
textbox(s,9.0,5.48,3.0,.56,'تفاوت اصلی: گستره، مرجع تصویب و ضمانت اجرا',11,GOLD,True)
#5
s=base(5,'هستهٔ درس','چهار دلیل اصلی نیاز جامعه به قانون')
reasons=[('۱','حفظ حقوق افراد','قانون اجازه نمی‌دهد زور، قدرت یا ثروت جای حق را بگیرد؛ فرد ضعیف نیز امکان دفاع دارد.'),('۲','نظم و امنیت','رفتارها پیش‌بینی‌پذیر می‌شوند و شهروندان با آرامش در جامعه زندگی می‌کنند.'),('۳','رشد مسئولیت‌پذیری','رعایت قانون، خودکنترلی، احترام و توجه به پیامد رفتار را تقویت می‌کند.'),('۴','حل عادلانهٔ اختلاف','دادگاه و مرجع رسمی، به‌جای دعوا و انتقام، با دلیل و قانون داوری می‌کنند.')]
for i,(n,h,b) in enumerate(reasons):
    x=.65+(i%2)*4.45; y=1.52+(i//2)*2.18; body_card(s,x,y,4.18,1.88,h,b,[GOLD,MINT,CORAL,ICE][i],15,11); textbox(s,x+3.45,y+.13,.46,.42,n,17,[GOLD,MINT,CORAL,ICE][i],True,'center')
add_art(s,'scale',9.3,1.43,3.3,4.65)
card(s,9.35,5.75,3.0,.6,fill=(75,53,24),line=GOLD); textbox(s,9.55,5.87,2.6,.32,'قانون = سپر حق + نقشهٔ نظم',11,GOLD,True,'center')
#6
s=base(6,'آزادی مسئولانه','آزادی واقعی؛ حق انتخاب همراه با توجه به پیامد')
body_card(s,.65,1.5,11.95,1.12,'تعریف تحلیلی','آزادی یعنی بتوانیم حق خود را اعمال کنیم؛ اما نه به قیمت آسیب به آرامش، امنیت، سلامت یا دارایی دیگران. مسئولیت، دشمن آزادی نیست؛ شرط پایدار ماندن آن است.',GOLD,17,14)
examples=[('خانه','گوش دادن به موسیقی حق من است؛ ولی صدای بلند در ساعت استراحت، حق آرامش خانواده را نقض می‌کند.',CORAL),('خیابان','رسیدن سریع مهم است؛ اما عبور بی‌محابا، جان خود و دیگران را در معرض خطر می‌گذارد.',GOLD),('مکان عمومی','دریافت خدمت حق همه است؛ رعایت صف، زمان و کرامت دیگران را حفظ می‌کند.',MINT)]
for i,(h,b,c) in enumerate(examples): body_card(s,.65+i*4.05,2.92,3.78,2.25,h,b,c,17,13)
card(s,2.15,5.55,9.0,.72,fill=(20,64,80),line=MINT); textbox(s,2.42,5.72,8.46,.35,'آزمون ساده: «اگر همه همین کار را انجام دهند، آیا نظم و حق دیگران باقی می‌ماند؟»',14,WHITE,True,'center')
#7
s=base(7,'سناریوی ۱','دنیای بدون قانون در رانندگی؛ یک زنجیرهٔ خطرناک')
steps=[('۱','نادیده‌گرفتن چراغ قرمز'),('۲','آشفتگی در تقاطع'),('۳','توقف جریان رفت‌وآمد'),('۴','تصادف، ترافیک و سلب امنیت')]
for i,(n,t) in enumerate(steps):
    y=1.55+i*1.02; card(s,.65,y,6.25,.78,accent=[CORAL,GOLD,MINT,ICE][i]); textbox(s,.95,y+.18,5.15,.34,t,14,WHITE,True); textbox(s,6.18,y+.17,.42,.38,n,15,[CORAL,GOLD,MINT,ICE][i],True,'center')
    if i<3: textbox(s,3.45,y+.76,.65,.24,'↓',15,GOLD,True,'center')
add_art(s,'traffic',7.15,1.35,5.25,4.95)
body_card(s,.65,5.72,11.7,.58,'نتیجهٔ تحلیلی','یک تخلف کوچک می‌تواند حقوق ده‌ها نفر را نقض کند؛ قانون رانندگی، محدودیت بی‌دلیل نیست، بلکه قرارداد جمعی برای حفظ جان و زمان است.',GOLD,12,10)
#8
s=base(8,'سناریوی ۲','اموال عمومی؛ دارایی مشترک امروز و فردای ما')
body_card(s,.65,1.5,5.2,3.85,'جامعهٔ باقانون','بوستان: پاکیزه و قابل استفاده برای همه\nاتوبوس: صندلی و تجهیزات سالم؛ رعایت نوبت\nکتابخانه: سکوت، امانت‌داری و دسترسی برابر\n\nپیامد: هزینهٔ کمتر، اعتماد بیشتر و کیفیت زندگی بالاتر. مسئولیت شهروندی یعنی چیزی را که متعلق به همه است، مثل دارایی خود حفظ کنیم.',MINT,18,14)
body_card(s,5.98,1.5,5.2,3.85,'جامعهٔ بی‌قانون','بوستان: تخریب فضای سبز و رهاکردن زباله\nاتوبوس: آسیب به تجهیزات و بی‌نظمی\nکتابخانه: سروصدا، گم‌شدن کتاب و محرومیت دیگران\n\nپیامد: منابع عمومی زود فرسوده می‌شوند و هزینهٔ بازسازی از پول همه پرداخت می‌شود؛ یعنی تخلف یک نفر، زیان همگانی است.',CORAL,18,14)
add_art(s,'shield',9.85,4.65,2.5,2.05)
card(s,1.15,5.65,8.55,.62,fill=(52,45,32),line=GOLD); textbox(s,1.38,5.8,8.1,.3,'سپر حقوق شهروندی: قانون + نظارت عمومی + وجدان مسئول',13,GOLD,True,'center')
#9
s=base(9,'از کوچک تا بزرگ','سه سطح قوانین و مقررات در زندگی روزمره')
levels=[('سطح ۱','گروهی و خانوادگی','تقسیم کار خانه، ساعت مطالعه، مقررات کلاس و اردو','نزدیک‌ترین محیط؛ توافق و همکاری مهم است.',MINT),('سطح ۲','شهری و عمومی','پارک، کتابخانه، اتوبوس، ورزشگاه و خیابان','افراد ناشناس نیز حق برابر دارند.',GOLD),('سطح ۳','کشوری و عمومی','قوانین راهنمایی، مالکیت، آموزش و حقوق شهروندی','قاعدهٔ رسمی برای همهٔ کشور و دارای ضمانت اجرا.',CORAL)]
for i,(n,h,b,c,col) in enumerate(levels):
    x=.65+i*4.02; y=1.55+i*.42; body_card(s,x,y,3.72,3.95,h,f'{b}\n\n{c}',col,18,14); textbox(s,x+.28,y+3.35,3.1,.35,n,13,col,True,'center')
textbox(s,1.25,6.25,10.8,.32,'هرچه گسترهٔ اثر یک رفتار بیشتر باشد، نیاز به قواعد رسمی‌تر و ضمانت اجرای قوی‌تر می‌شود.',13,ICE,True,'center')
#10
s=base(10,'فعالیت کارگاهی • صفحهٔ ۱۵','اردوی دانش‌آموزی؛ قانون‌گذاری برای یک تجربهٔ امن')
acts=[('چادر و وسایل','محل هر وسیله مشخص؛ وسایل مشترک با اجازه و سالم تحویل داده شود.'),('ساعات استراحت','خاموشی در ساعت توافق‌شده؛ سکوت برای حفظ حق استراحت همه.'),('آتش و ایمنی','روشن‌کردن آتش فقط در محل مجاز و با حضور مسئول؛ آب در دسترس باشد.'),('نوبت غذا','صف، سهم برابر، جلوگیری از اسراف و همکاری در جمع‌آوری.'),('کار گروهی','تقسیم نقش روشن؛ هرکس مسئول انجام وظیفه و گزارش نتیجه است.')]
for i,(h,b) in enumerate(acts):
    x=.65+(i%2)*4.15; y=1.45+(i//2)*1.55; body_card(s,x,y,3.9,1.33,h,b,[GOLD,MINT,CORAL,ICE,GOLD][i],14,10.5)
add_art(s,'camp',8.55,1.5,3.95,4.3)
card(s,8.72,5.72,3.45,.6,fill=(26,70,65),line=MINT); textbox(s,8.9,5.86,3.1,.31,'قاعدهٔ خوب: روشن، عادلانه و قابل اجرا',11,MINT,True,'center')
#11
s=base(11,'قانون متناسب با محیط','چرا هر مکان عمومی مقررات اختصاصی دارد؟')
places=[('کتابخانه','سکوت و امانت‌داری','زیرا تمرکز و دسترسی دیگران نباید مختل شود.',ICE),('پارک','پاکیزگی و حفظ فضای سبز','چون محیط و امکانات آن متعلق به همه و نسل آینده است.',MINT),('ورزشگاه','ایمنی، نوبت و احترام','برای پیشگیری از درگیری، ازدحام و آسیب جسمی.',CORAL),('خیابان','عبور قانونمند','چون یک تصمیم ناگهانی می‌تواند جان افراد زیادی را تهدید کند.',GOLD)]
for i,(h,r,b,c) in enumerate(places):
    x=.65+(i%2)*6.0; y=1.5+(i//2)*2.3; body_card(s,x,y,5.7,2.0,h,f'قاعدهٔ کلیدی: {r}\nدلیل: {b}',c,18,13)
body_card(s,2.35,6.02,8.65,.42,'قاعدهٔ طراحی مقررات','خطر محیط + حقوق افراد + هدف مکان = مقررات مناسب',GOLD,11,10)
#12
s=base(12,'دو راهنمای رفتار','قانون و اخلاق؛ هم‌پوشان، اما نه یکسان')
body_card(s,.65,1.5,5.0,3.75,'اخلاق','راهنمای رفتار درست، انسانی و فراتر از حداقل‌هاست. همیشه اجبار رسمی ندارد، اما وجدان و ارزش‌های انسانی پشتیبان آن‌اند.\n\nمثال: کمک داوطلبانه به فرد نابینا، گذشت، مهربانی و رعایت حال سالمندان.\n\nاخلاق می‌پرسد: «بهترین و انسانی‌ترین کار چیست؟»',CORAL,20,14)
body_card(s,5.82,1.5,5.0,3.75,'قانون','حداقل قاعدهٔ الزام‌آور برای حفظ نظم و حقوق عمومی است. ضمانت اجرای رسمی دارد و تخلف از آن ممکن است پیامد قانونی داشته باشد.\n\nمثال: ایستادن پشت چراغ قرمز، رعایت مالکیت و پرهیز از آسیب به اموال عمومی.\n\nقانون می‌پرسد: «چه کاری باید یا نباید انجام شود؟»',GOLD,20,14)
add_art(s,'ethics',9.65,4.55,2.75,2.15)
card(s,1.3,5.62,8.15,.72,fill=(25,59,70),line=MINT); textbox(s,1.55,5.8,7.65,.34,'جامعهٔ مطلوب فقط قانون‌مدار نیست؛ اخلاق‌مدار و مسئول نیز هست.',14,MINT,True,'center')
#13
s=base(13,'دانستنی تکمیلی','قانون اساسی؛ مادر همهٔ قوانین کشور')
body_card(s,.65,1.5,7.35,1.22,'تعریف','قانون اساسی، اصول بنیادین ادارهٔ کشور، ساختار نهادها و حقوق و مسئولیت‌های عمومی را مشخص می‌کند؛ به همین دلیل «مادر قوانین» نامیده می‌شود.',GOLD,17,13)
body_card(s,.65,2.92,7.35,1.22,'تصویب تاریخی','قانون اساسی جمهوری اسلامی ایران پس از انقلاب اسلامی، در همه‌پرسی ۱۲ آذر ۱۳۵۸ به رأی مردم گذاشته شد و تصویب شد.',MINT,17,13)
body_card(s,.65,4.34,7.35,1.42,'اصل برتری','قوانین عادی و مقررات پایین‌تر نباید با قانون اساسی مغایرت داشته باشند. این سلسله‌مراتب، از پراکندگی و تعارض قواعد جلوگیری می‌کند.',CORAL,17,13)
add_art(s,'doc',8.15,1.38,4.35,4.95)
textbox(s,8.7,6.03,3.2,.34,'قانون اساسی ← چارچوب مشترک حکومت و ملت',11,GOLD,True,'center')
#14
s=base(14,'تعادل شهروندی','حق و تکلیف؛ دو کفهٔ یک ترازوی اجتماعی')
body_card(s,.65,1.5,5.25,3.05,'حق','امتیاز و امکانی که فرد می‌تواند از آن برخوردار شود؛ مانند حق آموزش، امنیت، احترام، استفاده از امکانات عمومی و بیان نظر در چارچوب قانون.\n\nحق به ما می‌گوید: «چه چیزی باید برای من محترم شمرده شود؟»',MINT,21,14)
body_card(s,6.05,1.5,5.25,3.05,'تکلیف','مسئولیتی که فرد در برابر خود، دیگران و جامعه دارد؛ مانند رعایت مقررات مدرسه، حفظ اموال عمومی، احترام به نوبت و پرهیز از آسیب.\n\nتکلیف می‌گوید: «من برای حفظ حقوق چه کاری باید انجام دهم؟»',CORAL,21,14)
add_art(s,'scale',4.7,4.25,3.7,2.25)
textbox(s,.85,5.0,3.65,.75,'اگر فقط حق بخواهیم، حقوق دیگران نادیده گرفته می‌شود.',12,ICE,True,'center')
textbox(s,8.55,5.0,3.6,.75,'اگر تکلیف را انجام دهیم، حق همه پایدارتر می‌ماند.',12,GOLD,True,'center')
textbox(s,3.25,6.35,6.9,.3,'هیچ حقی بدون تکلیف پایدار نمی‌ماند؛ این دو، دو کفهٔ یک ترازو هستند.',13,MINT,True,'center')
#15
s=base(15,'سنجش کلاسی','پنج پرسش مرور درس • برای هر سؤال ۱۵ ثانیه')
qs=['چرا ورود به اجتماع باعث شکل‌گیری حقوق متقابل می‌شود؟','تفاوت اصلی «مقررات» و «قانون» چیست؟','چهار دلیل اصلی نیاز جامعه به قانون را نام ببرید.','در مثال چراغ قرمز، تخلف یک نفر چگونه به دیگران آسیب می‌زند؟','رابطهٔ حق و تکلیف را با یک مثال توضیح دهید.']
for i,q in enumerate(qs):
    y=1.36+i*1.0; card(s,.65,y,11.7,.78,accent=[GOLD,MINT,CORAL,ICE,GOLD][i]); textbox(s,1.0,y+.17,10.42,.38,q,13,WHITE,True); textbox(s,11.58,y+.15,.43,.4,str(i+1).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')),16,[GOLD,MINT,CORAL,ICE,GOLD][i],True,'center')
# timer bar
card(s,2.0,6.48,9.3,.28,fill=(17,49,69),line=(60,120,150));
for i in range(15):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(2.1+i*.6),Inches(6.56),Inches(.43),Inches(.08)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(*(GOLD if i<5 else MINT if i<10 else CORAL)); sh.line.fill.background()
textbox(s,5.8,6.79,1.8,.22,'۱۵ ثانیه فکر کن…',10,GOLD,True,'center')
#16
s=base(16,'پاسخنامه و جمع‌بندی','قانون یعنی احترام به حق خود و دیگران')
ans=[('۱','زیرا رابطه با دیگران، انتظارها و مسئولیت‌های دوطرفه ایجاد می‌کند.'),('۲','مقررات برای گروه یا محیط محدود است؛ قانون رسمی، عمومی و لازم‌الاجراست.'),('۳','حفظ حقوق، نظم و امنیت، رشد مسئولیت‌پذیری و حل عادلانهٔ اختلاف.'),('۴','آشفتگی تقاطع، ترافیک، تصادف و تهدید جان و زمان دیگران ایجاد می‌کند.'),('۵','حق استفاده از کتابخانه با تکلیف سکوت و امانت‌داری همراه است.')]
for i,(n,a) in enumerate(ans):
    y=1.35+i*.88; card(s,.65,y,8.4,.68,accent=[GOLD,MINT,CORAL,ICE,GOLD][i]); textbox(s,.95,y+.14,7.2,.37,a,11.5,WHITE,False); textbox(s,8.35,y+.12,.42,.4,n,14,[GOLD,MINT,CORAL,ICE,GOLD][i],True,'center')
card(s,9.35,1.35,3.0,2.05,fill=(70,49,22),line=GOLD); textbox(s,9.65,1.72,2.42,1.1,'«قانون یعنی\nاحترام به حق خود\nو دیگران.»',20,GOLD,True,'center')
card(s,9.35,3.65,3.0,2.45,fill=(18,65,77),line=MINT); textbox(s,9.65,3.94,2.42,.35,'سپاس از توجه شما',17,MINT,True,'center'); textbox(s,9.62,4.45,2.48,1.28,'گروه ۶ • کلاس ۷/۳\nامین کریم‌پور · عزیزی\nطاها عابدی‌فر · گله‌دار\nمحمدرضا فولادی',10.5,WHITE,False,'center')
textbox(s,1.25,6.18,7.2,.4,'نظم + امنیت + عدالت + مسئولیت = جامعهٔ قانون‌مدار',13,ICE,True,'center')

# transitions (fade / push / wipe), clickable and timed
for i,slide in enumerate(prs.slides):
    root=slide._element
    tr=OxmlElement('p:transition'); tr.set('spd','slow'); tr.set('advClick','1'); tr.set('advTm',str(11000 if i not in (14,15) else 20000))
    typ=['fade','push','wipe'][i%3]; child=OxmlElement(f'p:{typ}')
    if typ in ('push','wipe'): child.set('dir','l')
    tr.append(child)
    # place after clrMapOvr when possible
    idx=1
    for j,ch in enumerate(root):
        if ch.tag.endswith('clrMapOvr'): idx=j+1
    root.insert(idx,tr)

prs.core_properties.title='چرا به قانون نیاز داریم؟'
prs.core_properties.subject='درس ۳ مطالعات اجتماعی پایه هفتم'
prs.core_properties.author='گروه ۶ از کلاس ۷/۳'
prs.core_properties.keywords='قانون، مقررات، حقوق، تکالیف، مطالعات اجتماعی هفتم'
prs.save(OUT)

# create soft transition sounds and inject in package
def make_sound(path,kind):
    sr=22050; dur=.42 if kind!='impact' else .28; n=int(sr*dur); frames=[]
    random.seed({'whoosh':3,'impact':5,'shimmer':7}[kind])
    for i in range(n):
        t=i/sr; env=math.sin(math.pi*i/n)**2
        if kind=='whoosh': v=(random.random()*2-1)*env*(i/n)
        elif kind=='impact': v=(math.sin(2*math.pi*95*t)+.45*math.sin(2*math.pi*185*t))*math.exp(-12*t)
        else: v=(math.sin(2*math.pi*760*t)+.5*math.sin(2*math.pi*1140*t))*env*.45
        frames.append(struct.pack('<h',int(max(-1,min(1,v*.45))*32767)))
    with wave.open(path,'wb') as w: w.setparams((1,2,sr,n,'NONE','not compressed')); w.writeframes(b''.join(frames))
for k in ('whoosh','impact','shimmer'): make_sound(f'{AS}/{k}.wav',k)

# Sound injection into pptx (standards-based transition sound relationships)
import re, xml.etree.ElementTree as ET
PNS='http://schemas.openxmlformats.org/presentationml/2006/main'
RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
RELNS='http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('p',PNS); ET.register_namespace('a','http://schemas.openxmlformats.org/drawingml/2006/main'); ET.register_namespace('r',RNS)
tmp=OUT+'.tmp'
with zipfile.ZipFile(OUT,'r') as zin:
    entries={i.filename:zin.read(i.filename) for i in zin.infolist()}
    infos={i.filename:i for i in zin.infolist()}
for sn in range(1,17):
    relname=f'ppt/slides/_rels/slide{sn}.xml.rels'; kind=['whoosh','impact','shimmer'][(sn-1)%3]
    root=ET.fromstring(entries[relname]); ids=[int(x.attrib['Id'][3:]) for x in root if x.attrib.get('Id','').startswith('rId') and x.attrib['Id'][3:].isdigit()]; rid=f'rId{max(ids+[0])+1}'
    ET.SubElement(root,f'{{{RELNS}}}Relationship',{'Id':rid,'Type':'http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio','Target':f'../media/{kind}.wav'})
    entries[relname]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    sname=f'ppt/slides/slide{sn}.xml'; sroot=ET.fromstring(entries[sname]); tr=sroot.find(f'{{{PNS}}}transition')
    if tr is not None:
        sndAc=ET.SubElement(tr,f'{{{PNS}}}sndAc'); st=ET.SubElement(sndAc,f'{{{PNS}}}stSnd'); snd=ET.SubElement(st,f'{{{PNS}}}snd'); snd.set(f'{{{RNS}}}embed',rid); snd.set('name',kind)
    entries[sname]=ET.tostring(sroot,encoding='utf-8',xml_declaration=True)
ct=entries['[Content_Types].xml']
if b'Extension="wav"' not in ct: ct=ct.replace(b'</Types>',b'<Default Extension="wav" ContentType="audio/wav"/></Types>')
entries['[Content_Types].xml']=ct
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
    for name,data in entries.items(): zout.writestr(infos[name],data)
    for k in ('whoosh','impact','shimmer'): zout.write(f'{AS}/{k}.wav',f'ppt/media/{k}.wav')
os.replace(tmp,OUT)
print(OUT)
