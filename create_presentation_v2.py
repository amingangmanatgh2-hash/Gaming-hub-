from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFilter, ImageEnhance
import os, math, random, wave, struct, zipfile, xml.etree.ElementTree as ET

OUT='output/چرا_به_قانون_نیاز_داریم_نسخه_سینمایی_۳۰_اسلاید.pptx'
A='assets_v2'; os.makedirs(A,exist_ok=True); os.makedirs('output',exist_ok=True)
FONT='Vazirmatn'; NAVY=(4,16,30); DEEP=(7,27,49); ICE=(200,235,255); GOLD=(255,201,92); MINT=(78,225,186); CORAL=(255,113,102); WHITE=(244,250,255); BLUE=(68,175,235)
SW,SH=13.333,7.5; PW,PH=2560,1440
# Remove accidental English micro-labels from the generated classroom art while retaining its 3D scene.
school_raw=Image.open(f'{A}/hero_school.png').convert('RGB')
school_clean=school_raw.copy(); blurred=school_raw.filter(ImageFilter.GaussianBlur(24)); mask=Image.new('L',school_raw.size,0); md=ImageDraw.Draw(mask)
for box in [(500,285,650,370),(715,175,835,250),(900,280,1045,365),(500,515,690,615),(1000,515,1135,610)]: md.rounded_rectangle(box,18,fill=255)
school_clean.paste(blurred,(0,0),mask); school_clean.save(f'{A}/hero_school_clean.png')
heroes=[Image.open(f'{A}/hero_justice.png').convert('RGB'),school_clean,Image.open(f'{A}/hero_city.png').convert('RGB')]

def fit(im,w,h,zoom=1.0,dx=0,dy=0):
    scale=max(w/im.width,h/im.height)*zoom; nw,nh=int(im.width*scale),int(im.height*scale)
    x=max(0,min(nw-w,(nw-w)//2+dx)); y=max(0,min(nh-h,(nh-h)//2+dy))
    return im.resize((nw,nh),Image.Resampling.LANCZOS).crop((x,y,x+w,y+h))

def make_bg(i):
    random.seed(900+i); src=heroes[(i-1)%3]; base=fit(src,PW,PH,1.0+(i%4)*.025,(-1 if i%2 else 1)*(i%5)*22,(i%3-1)*18)
    base=ImageEnhance.Color(base).enhance(.82+(.08*(i%3))); base=ImageEnhance.Contrast(base).enhance(1.08)
    # cinematic tint and dark readable side
    over=Image.new('RGBA',(PW,PH),(2,10,20,25)); d=ImageDraw.Draw(over)
    left_dark = i%2==1
    for x in range(PW):
        t=x/(PW-1); a=int(208*((1-t)**1.7 if left_dark else t**1.7))
        d.line((x,0,x,PH),fill=(2,9,19,a))
    d.rectangle((0,0,PW,PH),outline=(120,220,255,35),width=4)
    # top and bottom vignette
    for y in range(230):
        a=int(80*(1-y/230)); d.line((0,y,PW,y),fill=(0,4,10,a)); d.line((0,PH-1-y,PW,PH-1-y),fill=(0,4,10,a))
    # fine luminous grid and particles
    for x in range(0,PW,160): d.line((x,0,x,PH),fill=(90,190,235,11),width=1)
    for y in range(0,PH,160): d.line((0,y,PW,y),fill=(90,190,235,9),width=1)
    for _ in range(110):
        x=random.randrange(PW); y=random.randrange(PH); r=random.choice((1,1,2,3)); col=random.choice((ICE,GOLD,MINT)); d.ellipse((x-r,y-r,x+r,y+r),fill=(*col,random.randrange(25,90)))
    img=Image.alpha_composite(base.convert('RGBA'),over).convert('RGB')
    # unique subtle grain increases visual texture and preserves high quality
    noise=Image.effect_noise((PW,PH),7+2*(i%4)).convert('L'); noise=ImageEnhance.Contrast(noise).enhance(.35)
    ng=Image.merge('RGB',(noise,noise,noise)); img=Image.blend(img,ng,.035)
    p=f'{A}/slide_bg_{i:02}.jpg'; img.save(p,'JPEG',quality=98,subsampling=0,optimize=False); return p

bgs=[make_bg(i) for i in range(1,31)]

# Animated transparent-look HUD overlays (GIFs animate during slide show)
def make_hud(path,variant):
    frames=[]; w,h=960,540
    for k in range(18):
        im=Image.new('RGBA',(w,h),(0,0,0,0)); d=ImageDraw.Draw(im)
        cx,cy=(760,150) if variant%2==0 else (190,370)
        for rr,c in [(105,GOLD),(145,ICE),(185,MINT)]:
            st=(k*13+rr)%360; d.arc((cx-rr,cy-rr,cx+rr,cy+rr),st,st+80+variant*9,fill=(*c,115),width=3)
        for j in range(26):
            ang=(j*0.71+k*.11+variant); rad=70+j*7; x=cx+math.cos(ang)*rad; y=cy+math.sin(ang)*rad*.55
            d.ellipse((x-2,y-2,x+2,y+2),fill=(*ICE,80+j*3))
        # moving scan line
        yy=(k*35+variant*61)%h; d.line((40,yy,w-40,yy),fill=(*BLUE,55),width=2)
        frames.append(im)
    frames[0].save(path,save_all=True,append_images=frames[1:],duration=70,loop=0,disposal=2,optimize=False)
for v in range(3): make_hud(f'{A}/hud_{v}.gif',v)

# richer stereo sound design
def sound(path,kind):
    sr=44100; dur={'whoosh':1.05,'impact':.72,'shimmer':1.25,'pulse':.82}[kind]; n=int(sr*dur); raw=[]; random.seed(40+len(kind))
    for i in range(n):
        t=i/sr; u=i/n
        if kind=='whoosh':
            env=(math.sin(math.pi*u)**1.7); noise=(random.random()*2-1); tone=math.sin(2*math.pi*(120+900*u)*t); v=.28*env*(.7*noise+.3*tone)
        elif kind=='impact':
            env=math.exp(-8.5*t); v=.48*env*(math.sin(2*math.pi*68*t)+.5*math.sin(2*math.pi*137*t))
        elif kind=='shimmer':
            env=math.sin(math.pi*u)**1.4; v=.20*env*(math.sin(2*math.pi*880*t)+.5*math.sin(2*math.pi*1320*t)+.25*math.sin(2*math.pi*1760*t))
        else:
            env=math.exp(-4*t); v=.32*env*math.sin(2*math.pi*(92+20*math.sin(t*7))*t)
        l=max(-1,min(1,v*(.9+.1*math.sin(t*5)))); r=max(-1,min(1,v*(.9+.1*math.cos(t*6))))
        raw.append(struct.pack('<hh',int(l*32767),int(r*32767)))
    with wave.open(path,'wb') as w: w.setparams((2,2,sr,n,'NONE','not compressed')); w.writeframes(b''.join(raw))
for k in ('whoosh','impact','shimmer','pulse'): sound(f'{A}/{k}.wav',k)

prs=Presentation(); prs.slide_width=Inches(SW); prs.slide_height=Inches(SH); blank=prs.slide_layouts[6]

def rtl(p,align=PP_ALIGN.RIGHT):
    p.alignment=align; pp=p._p.get_or_add_pPr(); pp.set('rtl','1'); pp.set('algn','r' if align==PP_ALIGN.RIGHT else 'ctr')

def text(sl,x,y,w,h,s,size=16,color=WHITE,bold=False,align='right',val=MSO_ANCHOR.MIDDLE,margin=.06):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=val
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(.025)
    p=tf.paragraphs[0]; p.text=s; rtl(p,PP_ALIGN.CENTER if align=='center' else PP_ALIGN.LEFT if align=='left' else PP_ALIGN.RIGHT)
    for r in p.runs:
        r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor(*color); r._r.get_or_add_rPr().set('lang','fa-IR')
    return sh

def shadow(sh,alpha=60,blur=70000,dist=36000):
    sp=sh._element.spPr; eff=sp.find(qn('a:effectLst'))
    if eff is None: eff=OxmlElement('a:effectLst'); sp.append(eff)
    o=OxmlElement('a:outerShdw'); o.set('blurRad',str(blur)); o.set('dist',str(dist)); o.set('dir','2700000'); o.set('algn','ctr'); o.set('rotWithShape','0')
    c=OxmlElement('a:srgbClr'); c.set('val','000000'); a=OxmlElement('a:alpha'); a.set('val',str(alpha*1000)); c.append(a); o.append(c); eff.append(o)

def panel(sl,x,y,w,h,fill=(7,25,44),line=BLUE,trans=10,accent=None):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(*fill); sh.fill.transparency=trans; sh.line.color.rgb=RGBColor(*line); sh.line.transparency=32; sh.line.width=Pt(1.4); shadow(sh)
    if accent:
        ac=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x+w-.055),Inches(y+.14),Inches(.045),Inches(h-.28)); ac.fill.solid(); ac.fill.fore_color.rgb=RGBColor(*accent); ac.line.fill.background()
    return sh

def pill(sl,x,y,w,label,col=GOLD):
    panel(sl,x,y,w,.36,fill=(17,47,66),line=col,trans=4); text(sl,x+.08,y+.03,w-.16,.27,label,10,col,True,'center')

def add_bg(sl,n):
    sl.shapes.add_picture(bgs[n-1],0,0,width=prs.slide_width,height=prs.slide_height)
    # subtle animated HUD overlay
    sl.shapes.add_picture(f'{A}/hud_{(n-1)%3}.gif',Inches(5.33 if n%2 else 0),Inches(0),width=Inches(8),height=Inches(4.5))

def head(sl,n,kicker,title):
    text(sl,.58,.28,12.15,.28,kicker,9.5,GOLD,True)
    text(sl,.58,.58,12.15,.58,title,25,WHITE,True)
    line=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(10.75),Inches(1.17),Inches(1.95),Inches(.025)); line.fill.solid(); line.fill.fore_color.rgb=RGBColor(*GOLD); line.line.fill.background()

def foot(sl,n):
    text(sl,.55,7.13,2.6,.18,'درس ۳ • مطالعات اجتماعی هفتم',8.5,(145,192,218),False,'left')
    fa=str(n).zfill(2).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))+' / ۳۰'
    text(sl,11.7,7.06,.95,.26,fa,10.5,GOLD,True)

def slide(n,kicker,title):
    s=prs.slides.add_slide(blank); add_bg(s,n); head(s,n,kicker,title); foot(s,n); return s

def card(sl,x,y,w,h,h1,b,col=GOLD,hs=15,bs=11.5):
    panel(sl,x,y,w,h,accent=col); text(sl,x+.2,y+.12,w-.4,.34,h1,hs,col,True); text(sl,x+.2,y+.54,w-.4,h-.66,b,bs,ICE,False,'right',MSO_ANCHOR.TOP)

def number(sl,x,y,n,col=GOLD):
    sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.5),Inches(.5)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(*col); sh.line.color.rgb=RGBColor(*WHITE); shadow(sh,45,40000,22000); text(sl,x,y+.02,.5,.43,str(n).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')),14,NAVY,True,'center')

def imgframe(sl,path,x,y,w,h):
    p=sl.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w),height=Inches(h)); shadow(p,70,90000,42000)
    fr=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x-.03),Inches(y-.03),Inches(w+.06),Inches(h+.06)); fr.fill.background(); fr.line.color.rgb=RGBColor(*GOLD); fr.line.transparency=25; fr.line.width=Pt(2); return p

# 01 cover
s=prs.slides.add_slide(blank); add_bg(s,1)
panel(s,.55,.55,6.2,5.65,fill=(3,15,29),line=GOLD,trans=9)
pill(s,.9,.88,2.65,'درس ۳ • مطالعات اجتماعی پایه هفتم')
text(s,.88,1.45,5.55,1.15,'چرا به قانون\nنیاز داریم؟',36,WHITE,True)
text(s,.9,2.82,5.3,.65,'قانون، مسیر رسیدن به نظم، امنیت و احترام به حقوق دیگران است.',17,ICE,False)
panel(s,.9,3.72,5.05,1.52,fill=(15,48,63),line=MINT,trans=6,accent=MINT)
text(s,1.12,3.92,4.6,.3,'گروه ۶ • کلاس ۷/۳',15,GOLD,True)
text(s,1.12,4.35,4.56,.66,'امین کریم‌پور · عزیزی · طاها عابدی‌فر\nگله‌دار · محمدرضا فولادی',12.5,WHITE,False)
text(s,.95,5.63,5.1,.28,'ارائهٔ سینمایی • نسخهٔ ۳۰ اسلایدی',11,MINT,True,'center'); foot(s,1)
#02
s=slide(2,'آغاز روایت','اگر فقط برای یک روز هیچ قانونی وجود نداشت…')
panel(s,.75,1.58,5.45,3.3,fill=(3,15,27),line=CORAL,trans=8)
text(s,1.08,1.94,4.75,.9,'چه کسی از حق ضعیف‌ترها دفاع می‌کرد؟',23,GOLD,True)
text(s,1.08,2.98,4.75,1.35,'خیابان، مدرسه، فروشگاه و حتی خانه خیلی زود دچار آشفتگی می‌شدند. قانون فقط مجموعه‌ای از «نبایدها» نیست؛ نقشه‌ای مشترک برای زندگی امن و عادلانه است.',15,ICE)
pill(s,1.15,5.18,4.55,'سؤال محوری: آزادی بدون مرز، آزادی است یا بی‌نظمی؟',MINT)
#03 roadmap
s=slide(3,'نقشهٔ مسیر','شش ایستگاه برای کشف معنای قانون')
road=[('۱','حق و اجتماع'),('۲','مقررات و قانون'),('۳','نظم و امنیت'),('۴','سناریوهای واقعی'),('۵','اخلاق و قانون اساسی'),('۶','مرور و چالش نهایی')]
for i,(n,t) in enumerate(road):
    x=.7+(i%3)*4.15; y=1.55+(i//3)*2.25; panel(s,x,y,3.85,1.65,fill=(8,35,55),line=[GOLD,MINT,CORAL,ICE,GOLD,MINT][i],trans=6); number(s,x+3.05,y+.25,n,[GOLD,MINT,CORAL,ICE,GOLD,MINT][i]); text(s,x+.25,y+.78,3.25,.4,t,16,WHITE,True,'center')
text(s,2.1,6.15,9.1,.34,'در هر ایستگاه: مفهوم ← مثال ← تحلیل ← نتیجهٔ کاربردی',13,GOLD,True,'center')
#04
s=slide(4,'گرم‌کردن کلاس','کدام رفتار، حق دیگران را نقض می‌کند؟')
opts=[('الف','پخش موسیقی با صدای بلند هنگام استراحت همسایه‌ها',CORAL),('ب','ایستادن در صف و رعایت نوبت',MINT),('پ','عبور موتور از پیاده‌رو برای فرار از ترافیک',CORAL),('ت','تحویل سالم کتاب امانتی',MINT)]
for i,(a,b,c) in enumerate(opts):
    x=.7+(i%2)*6.05; y=1.55+(i//2)*2.0; panel(s,x,y,5.7,1.55,fill=(10,34,52),line=c,trans=7,accent=c); text(s,x+.25,y+.3,4.85,.8,b,14,WHITE,True); pill(s,x+4.95,y+.98,.48,a,c)
pill(s,3.3,5.85,6.7,'پاسخ را فقط نگویید؛ دلیل و «حقِ نقض‌شده» را هم بیان کنید.',GOLD)
#05
s=slide(5,'مفهوم پایه','زندگی اجتماعی از کجا شروع می‌شود؟')
card(s,.7,1.48,5.55,2.05,'سه حلقهٔ نزدیک','خانواده: محبت، همکاری و احترام\nمدرسه: آموزش، امنیت و نظم\nمحله: همسایگی، رفت‌وآمد و امکانات مشترک',GOLD,18,13)
card(s,.7,3.75,5.55,2.0,'نکتهٔ تحلیلی','هرچه ارتباط ما با دیگران بیشتر شود، رفتار ما پیامدهای بیشتری پیدا می‌کند. بنابراین زندگی اجتماعی بدون توافق دربارهٔ حق‌ها و مسئولیت‌ها پایدار نمی‌ماند.',MINT,17,13)
imgframe(s,f'{A}/hero_school_clean.png',7.05,1.55,5.5,3.65); pill(s,7.55,5.5,4.5,'اجتماع = ارتباط + حق متقابل + مسئولیت',GOLD)
#06
s=slide(6,'شبکهٔ حقوق','حقوق متقابل؛ یک خیابان دوطرفه')
people=[('دانش‌آموز','حق آموزش و امنیت','تکلیف رعایت نظم'),('معلم','حق احترام و تمرکز کلاس','تکلیف آموزش عادلانه'),('خانواده','حق آرامش و اعتماد','تکلیف حمایت و مراقبت'),('شهروند','حق استفاده از شهر','تکلیف حفظ اموال عمومی')]
for i,(h,r,d) in enumerate(people):
    x=.65+(i%2)*6.08; y=1.45+(i//2)*2.22; card(s,x,y,5.76,1.88,h,f'{r}\n{d}',[GOLD,MINT,CORAL,ICE][i],17,12.5)
pill(s,2.45,6.08,8.45,'وقتی من تکلیفم را انجام می‌دهم، زمینهٔ برخورداری دیگری از حقش را می‌سازم.',GOLD)
#07
s=slide(7,'اصل طلایی','مرز حق من، حق دیگران و قانون است')
panel(s,.7,1.5,7.0,2.2,fill=(5,24,42),line=GOLD,trans=5)
text(s,1.05,1.85,6.3,1.42,'«هیچ‌کس در استفاده از حق خود، آزادی مطلق و بی‌قیدوشرط ندارد.»',25,GOLD,True,'center')
card(s,.7,4.0,3.35,1.75,'حق من','استفاده از امکانات، انتخاب و بیان نظر در چارچوب قانون',MINT,17,12)
card(s,4.23,4.0,3.35,1.75,'مرز مشترک','آرامش، امنیت، سلامت و دارایی دیگران',CORAL,17,12)
card(s,8.0,1.5,4.65,4.25,'آزمون سه‌سؤالی','۱. آیا رفتارم به کسی آسیب می‌زند؟\n۲. اگر همه این کار را کنند چه می‌شود؟\n۳. آیا حاضرم پیامدش را بپذیرم؟\n\nاگر پاسخ نگران‌کننده است، آزادی من از مرز مسئولیت عبور کرده است.',ICE,18,14)
#08
s=slide(8,'تعریف اول','مقررات؛ قواعد نزدیک برای یک محیط مشخص')
card(s,.7,1.5,5.5,3.85,'ویژگی‌های مقررات','• برای گروه یا محیط محدود تنظیم می‌شود.\n• هدف آن هماهنگی فعالیت‌های روزمره است.\n• می‌تواند با توافق اعضای گروه شکل بگیرد.\n• ضمانت اجرا معمولاً در همان محیط تعریف می‌شود.\n\nنمونه: ساعت خاموشی اردو، نوبت نظافت کلاس و زمان استفاده از تلفن همراه در خانه.',MINT,19,14)
imgframe(s,f'{A}/hero_school_clean.png',6.75,1.52,5.65,3.78); pill(s,7.35,5.65,4.5,'مقررات خوب: روشن، منصفانه و قابل اجرا',GOLD)
#09
s=slide(9,'تعریف دوم','قانون؛ قاعده‌ای رسمی، عمومی و لازم‌الاجرا')
card(s,.7,1.5,5.55,4.15,'چه چیزی قانون را متمایز می‌کند؟','• مرجع رسمی آن را تصویب می‌کند.\n• برای همهٔ افرادِ مشمول، اعتبار دارد.\n• هدف آن حفظ حقوق و نظم عمومی است.\n• نهادهای رسمی بر اجرای آن نظارت می‌کنند.\n• تخلف می‌تواند پیامد قانونی داشته باشد.\n\nمثال: قوانین راهنمایی‌ورانندگی، مالکیت و حفظ اموال عمومی.',GOLD,19,13.5)
imgframe(s,f'{A}/hero_justice.png',6.8,1.5,5.55,3.9); pill(s,7.3,5.68,4.55,'قانون، زبان مشترک جامعه برای حل اختلاف است.',MINT)
#10
s=slide(10,'مقایسهٔ دقیق','مقررات و قانون؛ شباهت‌ها و تفاوت‌ها')
cols=[('معیار','مقررات','قانون'),('گستره','گروه یا محیط محدود','جامعه یا بخش بزرگی از آن'),('مرجع','مدیر، خانواده یا توافق گروه','مرجع رسمی قانون‌گذاری'),('ضمانت اجرا','درون همان محیط','نهادهای رسمی و دادگاه'),('نمونه','قواعد کلاس و اردو','قوانین رانندگی و مالکیت')]
for r,row in enumerate(cols):
    y=1.42+r*.92; colors=[GOLD,MINT,ICE] if r==0 else [BLUE,MINT,GOLD]
    widths=[2.25,4.25,4.25]; xs=[.75,3.05,7.38]
    for c,val in enumerate(row): panel(s,xs[c],y,widths[c],.74,fill=(15,45,63) if r else (45,52,38),line=colors[c],trans=5); text(s,xs[c]+.12,y+.12,widths[c]-.24,.45,val,12.5 if r else 14,WHITE if r else colors[c],r==0,'center')
pill(s,2.45,6.23,8.6,'وجه مشترک: هر دو برای جلوگیری از آشفتگی و رعایت حقوق طراحی می‌شوند.',CORAL)
#11
s=slide(11,'هستهٔ درس','چهار دلیل اصلی نیاز به قانون')
reasons=[('حفظ حقوق','جلوگیری از زورگویی و دفاع از افراد',GOLD),('نظم و امنیت','پیش‌بینی‌پذیر شدن رفتارها و کاهش خطر',MINT),('مسئولیت‌پذیری','تقویت خودکنترلی و توجه به پیامد',CORAL),('حل اختلاف','داوری عادلانه به‌جای دعوا و انتقام',ICE)]
for i,(h,b,c) in enumerate(reasons):
    x=.7+(i%2)*6.05; y=1.48+(i//2)*2.15; panel(s,x,y,5.7,1.8,fill=(8,30,50),line=c,trans=5,accent=c); number(s,x+4.88,y+.24,i+1,c); text(s,x+.3,y+.28,4.3,.36,h,18,c,True); text(s,x+.3,y+.82,4.75,.55,b,12.5,ICE)
pill(s,2.8,6.05,7.75,'چهار پایهٔ یک جامعهٔ قانون‌مدار: حق، امنیت، مسئولیت و عدالت',GOLD)
#12
s=slide(12,'دلیل ۱','قانون، سپر افراد در برابر زورگویی است')
card(s,.7,1.5,5.7,2.05,'مسئله','اگر قدرت بدنی، ثروت یا نفوذ تعیین‌کننده باشد، فرد ضعیف راهی برای دفاع از حق خود ندارد و بی‌اعتمادی گسترش می‌یابد.',CORAL,18,14)
card(s,.7,3.78,5.7,2.0,'راه‌حل قانونی','قانون معیار مشترک می‌سازد؛ شکایت، ارائهٔ دلیل و رسیدگی رسمی را ممکن می‌کند و به همه—نه فقط قدرتمندان—حق دفاع می‌دهد.',MINT,18,14)
imgframe(s,f'{A}/hero_justice.png',7.15,1.55,5.15,3.45); pill(s,7.65,5.35,4.2,'عدالت یعنی معیار یکسان برای افراد متفاوت',GOLD)
#13
s=slide(13,'دلیل ۲','نظم و امنیت؛ نتیجهٔ رفتارهای پیش‌بینی‌پذیر')
steps=[('قاعدهٔ روشن','همه می‌دانند چه رفتاری انتظار می‌رود.'),('رعایت جمعی','رفتارها با یکدیگر هماهنگ می‌شوند.'),('کاهش خطر','احتمال درگیری، تصادف و اتلاف زمان کم می‌شود.'),('امنیت پایدار','افراد با اعتماد و آرامش فعالیت می‌کنند.')]
for i,(h,b) in enumerate(steps):
    x=.7+i*3.08; y=2.0+(i%2)*.38; panel(s,x,y,2.75,2.55,fill=(7,32,50),line=[GOLD,MINT,CORAL,ICE][i],trans=5); number(s,x+1.12,y-.3,i+1,[GOLD,MINT,CORAL,ICE][i]); text(s,x+.2,y+.45,2.35,.4,h,15,[GOLD,MINT,CORAL,ICE][i],True,'center'); text(s,x+.25,y+1.05,2.25,.95,b,11.5,ICE,False,'center')
text(s,2.15,5.85,9.1,.48,'امنیت فقط حضور پلیس نیست؛ نتیجهٔ انتخاب مسئولانهٔ میلیون‌ها شهروند است.',14,GOLD,True,'center')
#14
s=slide(14,'دلیل ۳','قانون‌مداری چگونه مسئولیت‌پذیری را رشد می‌دهد؟')
card(s,.7,1.45,5.5,4.35,'از اجبار بیرونی تا خودکنترلی','در آغاز ممکن است فرد فقط برای پرهیز از پیامد، قانون را رعایت کند. اما با فهم دلیل قانون و مشاهدهٔ اثر مثبت آن، رفتار مسئولانه به عادت تبدیل می‌شود.\n\nروند رشد:\nآگاهی از قاعده ← فهم پیامد ← انتخاب مسئولانه ← عادت مدنی\n\nقانون زمانی موفق‌تر است که شهروند «چرایی» آن را بداند، نه اینکه فقط از مجازات بترسد.',MINT,18,13.5)
card(s,6.72,1.45,5.6,1.75,'نمونهٔ مدرسه','دانش‌آموز پس از مدتی بدون تذکر، وسایل مشترک را سالم تحویل می‌دهد؛ چون می‌داند حق استفادهٔ دیگران نیز مهم است.',GOLD,17,12.5)
card(s,6.72,3.48,5.6,1.75,'نشانهٔ بلوغ اجتماعی','انجام کار درست حتی وقتی کسی ما را نمی‌بیند؛ یعنی تبدیل قانون بیرونی به وجدان مسئول درونی.',CORAL,17,12.5)
pill(s,7.2,5.62,4.7,'قانون‌مداری واقعی = فهم + انتخاب + مسئولیت',ICE)
#15
s=slide(15,'دلیل ۴','دادگاه؛ جایگزین عادلانهٔ دعوا و انتقام')
process=[('طرح اختلاف','هر طرف ادعای خود را بیان می‌کند.'),('بررسی دلیل','مدارک و گفته‌ها سنجیده می‌شود.'),('تطبیق با قانون','معیار تصمیم، قانون است نه قدرت.'),('رأی و اجرا','راه‌حل رسمی، روشن و قابل پیگیری ارائه می‌شود.')]
for i,(h,b) in enumerate(process):
    y=1.43+i*1.15; panel(s,.7,y,7.3,.9,fill=(8,31,50),line=[GOLD,MINT,CORAL,ICE][i],trans=6,accent=[GOLD,MINT,CORAL,ICE][i]); text(s,1.0,y+.14,1.75,.3,h,14,[GOLD,MINT,CORAL,ICE][i],True); text(s,2.7,y+.12,4.9,.38,b,11.5,WHITE); number(s,7.25,y+.19,i+1,[GOLD,MINT,CORAL,ICE][i])
imgframe(s,f'{A}/hero_justice.png',8.48,1.48,3.92,3.45); pill(s,8.72,5.28,3.45,'عدالت، آرامش را جایگزین انتقام می‌کند.',GOLD)
#16
s=slide(16,'سناریوی واقعی ۱','چراغ قرمز؛ چند ثانیه عجله، چندین حق نقض‌شده')
imgframe(s,f'{A}/hero_city.png',.7,1.48,6.5,4.4)
rights=[('جان و سلامت','خطر تصادف برای عابر و راننده'),('زمان','ایجاد ترافیک و تأخیر برای دیگران'),('آرامش','اضطراب و درگیری در تقاطع'),('اموال','خسارت به خودروها و امکانات شهری')]
for i,(h,b) in enumerate(rights):
    x=7.55+(i%2)*2.55; y=1.48+(i//2)*2.1; card(s,x,y,2.35,1.8,h,b,[CORAL,GOLD,MINT,ICE][i],14,10.5)
pill(s,2.0,6.18,9.25,'چراغ قرمز محدودیت بی‌دلیل نیست؛ قرارداد جمعی برای حفظ جان و زمان است.',GOLD)
#17
s=slide(17,'تحلیل زنجیره‌ای','یک تخلف چگونه به بحران تبدیل می‌شود؟')
chain=['نادیده‌گرفتن علامت','ورود هم‌زمان خودروها','قفل‌شدن تقاطع','ترافیک و تنش','تصادف و سلب امنیت']
for i,t in enumerate(chain):
    x=.55+i*2.55; y=2.25+(i%2)*.35; panel(s,x,y,2.25,1.42,fill=(12,36,51),line=[GOLD,MINT,CORAL,ICE,GOLD][i],trans=5); text(s,x+.18,y+.35,1.9,.55,t,12.5,WHITE,True,'center'); number(s,x+.86,y-.34,i+1,[GOLD,MINT,CORAL,ICE,GOLD][i])
    if i<4: text(s,x+2.2,y+.43,.4,.35,'←',17,GOLD,True,'center')
card(s,2.1,4.85,9.15,1.05,'درس سناریو','پیامد رفتار فقط به خود فرد محدود نمی‌ماند. جامعه شبکه‌ای به‌هم‌پیوسته است و قانون، از سرایت خطا جلوگیری می‌کند.',CORAL,15,12)
#18
s=slide(18,'سناریوی واقعی ۲','اموال عمومی؛ دارایی مشترک همهٔ ما')
card(s,.7,1.45,5.65,4.15,'جامعهٔ باقانون','بوستان پاکیزه و فضای سبز سالم\nاتوبوس و صندلی‌های قابل استفاده\nکتاب‌های سالم و دسترسی برابر\nهزینهٔ نگهداری کمتر\nاعتماد و کیفیت زندگی بیشتر\n\nشهروند مسئول، مال عمومی را مثل دارایی خود حفظ می‌کند؛ حتی دقیق‌تر، چون متعلق به همه است.',MINT,19,13.5)
card(s,6.48,1.45,5.65,4.15,'جامعهٔ بی‌قانون','زباله و تخریب فضای سبز\nآسیب به تجهیزات حمل‌ونقل\nگم‌شدن یا خراب‌شدن کتاب‌ها\nهزینهٔ بازسازی از پول همه\nمحرومیت نسل بعد از امکانات\n\nتخلف یک نفر ممکن است کوچک دیده شود، اما زیان آن میان هزاران نفر تقسیم می‌شود.',CORAL,19,13.5)
pill(s,2.35,6.0,8.7,'مال عمومی «مال هیچ‌کس» نیست؛ مال همه است و مسئولیت همه را می‌طلبد.',GOLD)
#19
s=slide(19,'قانون متناسب با محیط','چرا هر مکان عمومی، مقررات مخصوص دارد؟')
places=[('کتابخانه','سکوت و امانت‌داری','تمرکز و دسترسی برابر',ICE),('پارک','پاکیزگی و حفظ گیاهان','سلامت و محیط زیست',MINT),('ورزشگاه','ایمنی و احترام','پیشگیری از ازدحام و درگیری',CORAL),('خیابان','عبور قانونمند','حفظ جان و جریان رفت‌وآمد',GOLD)]
for i,(h,r,why,c) in enumerate(places):
    x=.7+(i%2)*6.05; y=1.45+(i//2)*2.15; card(s,x,y,5.7,1.85,h,f'قاعده: {r}\nدلیل: {why}',c,17,12)
pill(s,2.75,6.0,7.9,'فرمول مقررات مناسب: هدف مکان + خطرها + حقوق افراد',GOLD)
#20
s=slide(20,'سه سطح قواعد','از خانه تا کشور؛ هر سطح چه نیازی دارد؟')
levels=[('سطح ۱','خانوادگی و گروهی','تقسیم کار خانه، مقررات کلاس و اردو','توافق و همکاری',MINT),('سطح ۲','شهری و عمومی','پارک، اتوبوس، ورزشگاه و خیابان','حقوق افراد ناشناس',GOLD),('سطح ۳','کشوری و رسمی','مالکیت، آموزش و حقوق شهروندی','قانون‌گذاری و ضمانت اجرا',CORAL)]
for i,(n,h,b,k,c) in enumerate(levels):
    x=.65+i*4.16; y=1.62+i*.25; panel(s,x,y,3.85,3.9,fill=(8,31,48),line=c,trans=5,accent=c); pill(s,x+.55,y+.32,2.75,n,c); text(s,x+.3,y+1.02,3.2,.42,h,17,c,True,'center'); text(s,x+.35,y+1.72,3.1,.85,b,12.5,WHITE,False,'center'); text(s,x+.35,y+2.85,3.1,.45,k,11.5,ICE,True,'center')
text(s,1.65,6.15,10.0,.35,'هرچه گسترهٔ اثر رفتار بیشتر باشد، قواعد رسمی‌تر و ضمانت اجرا قوی‌تر می‌شود.',13,GOLD,True,'center')
#21
s=slide(21,'کارگاه صفحهٔ ۱۵','اردوی دانش‌آموزی؛ شما قانون‌گذار هستید')
imgframe(s,f'{A}/hero_city.png',.7,1.5,6.75,4.45)
card(s,7.8,1.5,4.55,1.35,'ماموریت','برای اردو مقرراتی بنویسید که هم امنیت را حفظ کند، هم عادلانه باشد و هم واقعاً قابل اجرا باشد.',GOLD,16,12)
card(s,7.8,3.05,4.55,1.2,'معیار ۱: روشنی','آیا همه دقیقاً می‌فهمند چه کاری باید انجام دهند؟',MINT,15,11)
card(s,7.8,4.45,4.55,1.2,'معیار ۲ و ۳','آیا قاعده منصفانه و اجرای آن در شرایط اردو ممکن است؟',CORAL,15,11)
pill(s,2.25,6.2,8.8,'قانون خوب فقط سخت‌گیرانه نیست؛ هدفمند، روشن و قابل دفاع است.',ICE)
#22
s=slide(22,'کارگاه اردو','پنج موقعیت، پنج تصمیم مسئولانه')
camps=[('چادر و وسایل','جای مشخص؛ استفاده با اجازه؛ تحویل سالم',GOLD),('ساعات استراحت','خاموشی توافقی و رعایت سکوت',MINT),('آتش و ایمنی','محل مجاز، حضور مسئول و آب آماده',CORAL),('نوبت غذا','صف، سهم برابر و جلوگیری از اسراف',ICE),('کار گروهی','تقسیم نقش، زمان‌بندی و گزارش نتیجه',GOLD)]
for i,(h,b,c) in enumerate(camps):
    x=.7+(i%2)*4.25 if i<4 else 4.85; y=1.42+(i//2)*1.55 if i<4 else 4.62; card(s,x,y,4.0,1.3,h,b,c,14,10.5)
imgframe(s,f'{A}/hero_city.png',9.18,1.46,3.1,3.1); pill(s,9.25,4.85,2.95,'امنیت اردو = همکاری همه',MINT)
#23
s=slide(23,'تمرین تصمیم‌گیری','آیا این مقرره، خوب طراحی شده است؟')
card(s,.7,1.48,5.55,2.0,'پیشنهاد','«هیچ‌کس در اردو حق صحبت‌کردن ندارد.»\n\nآیا این قاعده هدفمند، منصفانه و متناسب است؟ چه تغییری پیشنهاد می‌کنید؟',CORAL,18,14)
card(s,.7,3.75,5.55,2.05,'بازنویسی بهتر','«از ساعت ۲۲ تا ۷ صبح در محدودهٔ چادرها سکوت رعایت شود؛ گفت‌وگوی ضروری با صدای آهسته و در محل مشخص انجام شود.»',MINT,18,14)
criteria=[('روشن؟','زمان و مکان مشخص دارد.'),('عادلانه؟','حق استراحت و ارتباط را متعادل می‌کند.'),('قابل اجرا؟','رفتار مورد انتظار قابل مشاهده است.')]
for i,(h,b) in enumerate(criteria): card(s,6.72,1.5+i*1.55,5.55,1.3,h,b,[GOLD,MINT,ICE][i],15,11.5)
pill(s,7.2,6.12,4.6,'قانون خوب، مسئله را حل می‌کند؛ نه اینکه مسئلهٔ تازه بسازد.',GOLD)
#24
s=slide(24,'دو راهنمای رفتار','قانون و اخلاق؛ تفاوت و هم‌پوشانی')
card(s,.7,1.45,5.65,4.25,'اخلاق','راهنمای رفتار درست و انسانی است و معمولاً وجدان، ارزش‌ها و قضاوت اجتماعی پشتیبان آن‌اند.\n\nنمونه‌ها:\nکمک به فرد نابینا، گذشت، مهربانی و رعایت حال سالمندان.\n\nپرسش اخلاق: «بهترین و انسانی‌ترین کار چیست؟»',CORAL,20,14)
card(s,6.48,1.45,5.65,4.25,'قانون','حداقل قاعدهٔ الزام‌آور برای حفظ نظم و حقوق عمومی است و ضمانت اجرای رسمی دارد.\n\nنمونه‌ها:\nایستادن پشت چراغ قرمز، رعایت مالکیت و حفظ اموال عمومی.\n\nپرسش قانون: «چه کاری باید یا نباید انجام شود؟»',GOLD,20,14)
pill(s,2.2,6.07,8.95,'جامعهٔ مطلوب فقط قانون‌مدار نیست؛ اخلاق‌مدار و مسئول نیز هست.',MINT)
#25
s=slide(25,'دانستنی تکمیلی','قانون اساسی؛ مادر همهٔ قوانین کشور')
imgframe(s,f'{A}/hero_justice.png',.7,1.48,5.1,4.25)
card(s,6.2,1.48,6.05,1.4,'تعریف','اصول بنیادین ادارهٔ کشور، ساختار نهادها و حقوق و مسئولیت‌های عمومی را مشخص می‌کند.',GOLD,17,12.5)
card(s,6.2,3.08,6.05,1.35,'تصویب تاریخی','پس از انقلاب اسلامی، در همه‌پرسی ۱۲ آذر ۱۳۵۸ به رأی مردم گذاشته شد و تصویب شد.',MINT,17,12.5)
card(s,6.2,4.63,6.05,1.35,'اصل برتری','قوانین و مقررات دیگر نباید با قانون اساسی مغایرت داشته باشند.',CORAL,17,12.5)
pill(s,2.5,6.1,8.4,'قانون اساسی، چارچوب مشترک رابطهٔ حکومت، نهادها و ملت است.',ICE)
#26
s=slide(26,'تعادل شهروندی','حق و تکلیف؛ دو کفهٔ جدانشدنی')
card(s,.7,1.45,5.55,3.8,'حق','امتیاز و امکانی که باید برای فرد محترم شمرده شود؛ مانند حق آموزش، امنیت، احترام، استفاده از امکانات عمومی و بیان نظر در چارچوب قانون.\n\nحق می‌پرسد: «چه چیزی باید برای من فراهم و محترم باشد؟»',MINT,21,14)
card(s,6.78,1.45,5.55,3.8,'تکلیف','مسئولیتی که فرد در برابر خود، دیگران و جامعه دارد؛ مانند رعایت مقررات، حفظ اموال عمومی، احترام به نوبت و پرهیز از آسیب.\n\nتکلیف می‌پرسد: «من برای حفظ حقوق چه باید بکنم؟»',CORAL,21,14)
panel(s,3.05,5.55,7.2,.68,fill=(47,48,32),line=GOLD,trans=3); text(s,3.3,5.71,6.7,.34,'هیچ حقی بدون تکلیف پایدار نمی‌ماند؛ تعادل، شرط عدالت است.',14,GOLD,True,'center')
#27
s=slide(27,'سناریوی تکمیلی','قانون در فضای مجازی و گروه کلاسی')
card(s,.7,1.45,5.55,4.3,'مسئلهٔ تازه، همان حقوق قدیمی','انتشار عکس هم‌کلاسی بدون اجازه، ارسال پیام در ساعت نامناسب، تمسخر در گروه و پخش خبر تأییدنشده ممکن است در فضای مجازی رخ دهد؛ اما حقوق افراد واقعی‌اند.\n\nحقوق در خطر: حریم خصوصی، آبرو، آرامش و دسترسی برابر به اطلاعات درست.',CORAL,18,13.5)
card(s,6.72,1.45,5.6,4.3,'منشور پیشنهادی گروه کلاس','۱. انتشار تصویر فقط با اجازه\n۲. پیام در ساعت توافق‌شده\n۳. نقد محترمانه، بدون تمسخر\n۴. بررسی منبع پیش از بازنشر\n۵. گزارش پیام آسیب‌زا به مسئول گروه\n\nقانون‌مداری، در صفحهٔ نمایش متوقف نمی‌شود.',MINT,18,13)
pill(s,2.6,6.15,8.15,'قبل از ارسال: درست است؟ ضروری است؟ محترمانه است؟',GOLD)
#28
s=slide(28,'چالش نهایی','پنج پرسش برای سنجش کلاس • هر سؤال ۱۵ ثانیه')
qs=['چرا زندگی اجتماعی، حقوق متقابل ایجاد می‌کند؟','سه تفاوت مهم مقررات و قانون چیست؟','چهار دلیل اصلی نیاز جامعه به قانون را نام ببرید.','در سناریوی چراغ قرمز، کدام حقوق نقض می‌شود؟','رابطهٔ حق، تکلیف، قانون و اخلاق را توضیح دهید.']
for i,q in enumerate(qs):
    y=1.33+i*1.0; panel(s,.7,y,11.85,.78,fill=(8,31,49),line=[GOLD,MINT,CORAL,ICE,GOLD][i],trans=4,accent=[GOLD,MINT,CORAL,ICE,GOLD][i]); number(s,11.7,y+.14,i+1,[GOLD,MINT,CORAL,ICE,GOLD][i]); text(s,1.0,y+.12,10.45,.42,q,12.5,WHITE,True)
# timer
for i in range(15):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(2.0+i*.6),Inches(6.55),Inches(.45),Inches(.1)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(*(GOLD if i<5 else MINT if i<10 else CORAL)); sh.line.fill.background()
text(s,5.7,6.72,1.95,.23,'۱۵ ثانیه فکر کن…',10,GOLD,True,'center')
#29
s=slide(29,'پاسخنامه','پاسخ‌های کوتاه، دقیق و تشریحی')
answers=[('۱','ارتباط با دیگران، انتظار و مسئولیت دوطرفه ایجاد می‌کند.'),('۲','گستره، مرجع تصویب و نوع ضمانت اجرا متفاوت است.'),('۳','حفظ حقوق، نظم و امنیت، مسئولیت‌پذیری و حل عادلانهٔ اختلاف.'),('۴','جان، زمان، آرامش و دارایی افراد در معرض آسیب قرار می‌گیرد.'),('۵','حق با تکلیف پایدار می‌شود؛ قانون حداقل الزام و اخلاق رفتار برتر را نشان می‌دهد.')]
for i,(n,a) in enumerate(answers):
    y=1.33+i*1.0; panel(s,.7,y,11.85,.8,fill=(8,31,49),line=[GOLD,MINT,CORAL,ICE,GOLD][i],trans=4,accent=[GOLD,MINT,CORAL,ICE,GOLD][i]); pill(s,11.62,y+.2,.55,n,[GOLD,MINT,CORAL,ICE,GOLD][i]); text(s,1.0,y+.14,10.35,.42,a,12,WHITE,False)
pill(s,2.4,6.55,8.6,'پاسخ خوب فقط «چیست» را نمی‌گوید؛ «چرا» و «مثال» هم دارد.',MINT)
#30 finale
s=prs.slides.add_slide(blank); add_bg(s,30)
panel(s,.7,.72,6.2,5.8,fill=(3,15,28),line=GOLD,trans=7)
pill(s,1.05,1.02,2.45,'جمع‌بندی نهایی',GOLD)
text(s,1.0,1.72,5.55,1.25,'قانون یعنی احترام\nبه حق خود و دیگران.',30,GOLD,True)
text(s,1.03,3.23,5.45,.85,'حق بدون تکلیف ناپایدار است؛ آزادی بدون مسئولیت به بی‌نظمی می‌رسد؛ و جامعه بدون عدالت، اعتماد خود را از دست می‌دهد.',15,ICE)
panel(s,1.02,4.38,5.2,1.25,fill=(17,57,68),line=MINT,trans=4)
text(s,1.25,4.57,4.72,.28,'سپاس از توجه شما',17,MINT,True,'center')
text(s,1.24,4.94,4.74,.48,'امین کریم‌پور · عزیزی · طاها عابدی‌فر\nگله‌دار · محمدرضا فولادی',11.5,WHITE,False,'center')
text(s,1.28,5.93,4.65,.26,'گروه ۶ • کلاس ۷/۳',11,GOLD,True,'center'); foot(s,30)

# transitions: varied cinematic + auto/click
for i,sl in enumerate(prs.slides):
    root=sl._element; tr=OxmlElement('p:transition'); tr.set('spd','slow'); tr.set('advClick','1'); tr.set('advTm',str(11500 if i not in (27,28,29) else 20000))
    typ=['fade','push','wipe','cover'][i%4]; ch=OxmlElement(f'p:{typ}');
    if typ in ('push','wipe','cover'): ch.set('dir',['l','r','u','d'][i%4])
    tr.append(ch); idx=1
    for j,e in enumerate(root):
        if e.tag.endswith('clrMapOvr'): idx=j+1
    root.insert(idx,tr)

prs.core_properties.title='چرا به قانون نیاز داریم؟ — نسخه سینمایی ۳۰ اسلایدی'
prs.core_properties.subject='درس ۳ مطالعات اجتماعی پایه هفتم'
prs.core_properties.author='گروه ۶ از کلاس ۷/۳'
prs.core_properties.comments='ارائهٔ کلاس و کنفرانس با تصاویر سه‌بعدی، افکت‌های متحرک و صدای سینمایی'
prs.save(OUT)

# Embed transition sounds and relationships
PNS='http://schemas.openxmlformats.org/presentationml/2006/main'; RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'; RELNS='http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('p',PNS); ET.register_namespace('a','http://schemas.openxmlformats.org/drawingml/2006/main'); ET.register_namespace('r',RNS)
tmp=OUT+'.tmp'
with zipfile.ZipFile(OUT) as zin:
    info={z.filename:z for z in zin.infolist()}; data={z.filename:zin.read(z.filename) for z in zin.infolist()}
for sn in range(1,31):
    rel=f'ppt/slides/_rels/slide{sn}.xml.rels'; kind=['whoosh','impact','shimmer','pulse'][(sn-1)%4]; rr=ET.fromstring(data[rel]); ids=[int(x.attrib['Id'][3:]) for x in rr if x.attrib.get('Id','')[3:].isdigit()]; rid=f'rId{max(ids+[0])+1}'
    ET.SubElement(rr,f'{{{RELNS}}}Relationship',{'Id':rid,'Type':'http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio','Target':f'../media/sfx_{kind}.wav'}); data[rel]=ET.tostring(rr,encoding='utf-8',xml_declaration=True)
    sx=f'ppt/slides/slide{sn}.xml'; sr=ET.fromstring(data[sx]); tr=sr.find(f'{{{PNS}}}transition')
    sndAc=ET.SubElement(tr,f'{{{PNS}}}sndAc'); st=ET.SubElement(sndAc,f'{{{PNS}}}stSnd'); snd=ET.SubElement(st,f'{{{PNS}}}snd'); snd.set(f'{{{RNS}}}embed',rid); snd.set('name',kind); data[sx]=ET.tostring(sr,encoding='utf-8',xml_declaration=True)
ct=data['[Content_Types].xml']
if b'Extension="wav"' not in ct: ct=ct.replace(b'</Types>',b'<Default Extension="wav" ContentType="audio/wav"/></Types>')
data['[Content_Types].xml']=ct
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as zout:
    for n,b in data.items(): zout.writestr(info[n],b)
    for k in ('whoosh','impact','shimmer','pulse'): zout.write(f'{A}/{k}.wav',f'ppt/media/sfx_{k}.wav')
os.replace(tmp,OUT)
print(OUT,os.path.getsize(OUT)/1024/1024,'MB')
